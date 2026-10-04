"""Manual release log -> repeater kitting counter (no SAP needed)."""
import sys, json, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule, CellIsRule, DataBarRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter as CL
from openpyxl.workbook.properties import CalcProperties

OUT=sys.argv[1]; CFG=json.load(open(sys.argv[2])); TEST=json.load(open(sys.argv[3])) if len(sys.argv)>3 else None
PLAN=CFG["plan"]; NP=len(PLAN); NREP=CFG["repeaters"]; D=CFG["details"]
NAVY="141A26"; INK="1E2433"; ROSE="B0245E"; GREY="5D6576"; LINE="D9DDE5"; PANEL="F3F4F7"; WHITE="FFFFFF"; MUTE="AEB6C6"
OK="1F8A5B"; OK_S="DDF2E7"; WARN="A8660B"; WARN_S="FBF0DC"; CRIT="C0392B"; CRIT_S="FBE6E3"; INFO="3558B8"; INFO_S="E3EAFA"; BLUE="0000FF"; YEL="FFF2B3"
F="Arial"
def font(sz=11,b=False,c=INK,i=False): return Font(name=F,size=sz,bold=b,color=c,italic=i)
def fill(c): return PatternFill("solid",start_color=c,end_color=c)
thin=Side(style="thin",color=LINE); BOX=Border(top=thin,bottom=thin,left=thin,right=thin); BOT=Border(bottom=thin)
C=Alignment(horizontal="center",vertical="center",wrap_text=True); L=Alignment(horizontal="left",vertical="center",wrap_text=True)
def hdr(ws,r,c,t,bg=INK):
    x=ws.cell(r,c,t); x.font=font(10,True,WHITE); x.fill=fill(bg); x.alignment=C; x.border=BOX

wb=Workbook()
# ---------------- Contract (plan) ----------------
kp=wb.active; kp.title="Contract"; kp.sheet_view.showGridLines=False
kp["A1"]="CONTRACT — WHAT ONE REPEATER NEEDS"; kp["A1"].font=font(16,True,NAVY)
kp["A2"]="Blue cells are inputs. Change these only when the contract or kitting plan changes."; kp["A2"].font=font(10,c=GREY,i=True)
info=[("Contract",D["contract"]),("Program",D["program"]),("ERP item",D["erp"]),("Config",D["config"]),("Contract repeaters",NREP)]
for i,(k,v) in enumerate(info):
    kp.cell(3+i,1,k).font=font(10,True); c=kp.cell(3+i,2,v); c.font=font(11,True,BLUE); c.fill=fill(YEL)
P0=10; P1=P0+NP-1
for j,h in enumerate(["Sequence","Assy Code","Assy Description","Qty per Repeater","Opening Balance","Contract Qty"],1): hdr(kp,P0-1,j,h,NAVY)
kp["E8"]="Already released before this log started (leave 0 to start fresh)"; kp["E8"].font=font(8,c=GREY,i=True)
for i,(s,code,d,per) in enumerate(PLAN):
    r=P0+i
    for j,v in enumerate([s,code,d,per,0],1):
        x=kp.cell(r,j,v); x.font=font(10,c=BLUE); x.border=BOX; x.alignment=L if j<4 else C
    x=kp.cell(r,6,f"=D{r}*$B$7"); x.font=font(10); x.border=BOX; x.alignment=C
for c,w in zip("ABCDEF",[20,18,46,16,16,14]): kp.column_dimensions[c].width=w
K_CODE=f"Contract!$B${P0}:$B${P1}"; K_DESC=f"Contract!$C${P0}:$C${P1}"; NREPS="Contract!$B$7"

# ---------------- Release Log ----------------
lg=wb.create_sheet("Release Log"); lg.sheet_view.showGridLines=False
R0,R1=6,6+CFG.get("rows",1500)-1
lg["A1"]="RELEASE LOG"; lg["A1"].font=font(16,True,NAVY)
lg["A2"]="One line per release to production. Type the date, pick the assy code, enter the qty released. The Dashboard counts repeaters from this."; lg["A2"].font=font(10,c=GREY,i=True)
lg["A3"]="Example:  03-Oct-26 · 92RRA00548AAA · 24 · N.Bakare · for Rep 14"; lg["A3"].font=font(10,c=GREY)
for j,h in enumerate(["Date","Assy Code","Description","Qty Released","Released By","Note (WO / repeater)","Check"],1): hdr(lg,5,j,h,ROSE if j in (1,2,4,5,6) else NAVY)
for r in range(R0,R1+1):
    lg[f"C{r}"]=f'=IF(B{r}="","",IFERROR(INDEX({K_DESC},MATCH(B{r},{K_CODE},0)),"Unknown code"))'
    lg[f"G{r}"]=f'=IF(AND(B{r}="",D{r}=""),"",IF(B{r}="","Missing code",IF(C{r}="Unknown code","Unknown code",IF(D{r}="","Missing qty",IF(A{r}="","Missing date","")))))'
    for j in range(1,8):
        x=lg.cell(r,j); x.border=BOX; x.font=font(10,c=BLUE if j in (1,2,4,5,6) else GREY); x.alignment=C if j in (1,4,7) else L
    lg[f"A{r}"].number_format="dd-mmm-yy"; lg[f"D{r}"].font=font(11,True,BLUE)
for c,w in zip("ABCDEFG",[12,18,42,13,16,26,14]): lg.column_dimensions[c].width=w
lg.freeze_panes="A6"; lg.auto_filter.ref=f"A5:G{R1}"
dv=DataValidation(type="list",formula1=f"=Contract!$B${P0}:$B${P1}",allow_blank=True,showErrorMessage=True,error="Pick an assy code from the contract."); lg.add_data_validation(dv); dv.add(f"B{R0}:B{R1}")
dq=DataValidation(type="whole",operator="notEqual",formula1="0",allow_blank=True,showErrorMessage=True,error="Enter a whole number (use a negative number to correct a mistake)."); lg.add_data_validation(dq); dq.add(f"D{R0}:D{R1}")
dd=DataValidation(type="date",operator="greaterThan",formula1="40000",allow_blank=True,showErrorMessage=True,error="Enter a date, e.g. 03/10/2026"); lg.add_data_validation(dd); dd.add(f"A{R0}:A{R1}")
lg.conditional_formatting.add(f"G{R0}:G{R1}",FormulaRule(formula=[f'$G{R0}<>""'],fill=fill(CRIT_S),font=Font(name=F,color=CRIT,bold=True)))
LB=f"'Release Log'!$B${R0}:$B${R1}"; LD=f"'Release Log'!$D${R0}:$D${R1}"; LA=f"'Release Log'!$A${R0}:$A${R1}"
if TEST:
    for i,(dt,code,q,who,note) in enumerate(TEST):
        r=R0+i; lg[f"A{r}"]=datetime.date.fromisoformat(dt); lg[f"B{r}"]=code; lg[f"D{r}"]=q; lg[f"E{r}"]=who; lg[f"F{r}"]=note

# ---------------- Dashboard ----------------
ds=wb.create_sheet("Dashboard",0); ds.sheet_view.showGridLines=False; ds.sheet_view.zoomScale=80
cols={"A":2,"B":14,"C":18,"D":38,"E":10,"F":12,"G":11,"H":12,"I":13,"J":15,"K":2,"L":24}
for k,v in cols.items(): ds.column_dimensions[k].width=v
EDGE="J"
for c in "BCDEFGHIJ":
    for r in (1,2,3): ds[f"{c}{r}"].fill=fill(NAVY)
ds.row_dimensions[2].height=40
ds["B2"]=CFG["title"]; ds["B2"].font=font(24,True,WHITE); ds["B2"].alignment=Alignment(vertical="center")
ds["B3"]='="Contract "&Contract!B3&"  ·  "&Contract!B4&"  ·  "&Contract!B5&"  ·  "&Contract!B6&"  ·  Contract: "&Contract!B7&" repeaters  ·  "&COUNT(' + LD + ')&" releases logged"'
ds["B3"].font=font(11,c=MUTE)
# request input
ds["I2"]="REPEATERS REQUESTED ▸"; ds["I2"].font=font(10,True,MUTE); ds["I2"].alignment=Alignment(horizontal="right",vertical="center",wrap_text=True)
ds["J2"]=CFG.get("request",2); ds["J2"].font=font(26,True,WHITE); ds["J2"].fill=fill(ROSE); ds["J2"].alignment=C
ds["J3"]="type the target here"; ds["J3"].font=font(8,c=MUTE,i=True); ds["J3"].alignment=C
dr=DataValidation(type="whole",operator="between",formula1="0",formula2="999",showErrorMessage=True,error="Enter a whole number of repeaters"); ds.add_data_validation(dr); dr.add("J2")
REQ="$J$2"
T0=12; T1=T0+NP-1
# tiles rows 5-7
tiles=[("B","C","REPEATERS KITTED",f"=MIN(G{T0}:G{T1})",f'="of "&{NREPS}&" in contract · counts when all {NP} codes are released"'),
       ("D","D","REQUEST STATUS",f'=IF(MIN(G{T0}:G{T1})>={REQ},"READY",MAX(0,{REQ}-MIN(G{T0}:G{T1}))&" TO GO")',f'="target "&{REQ}&" repeaters · "&COUNTIF(J{T0}:J{T1},"READY")&" of {NP} codes ready"'),
       ("E","G","UNITS TO RELEASE",f"=SUM(I{T0}:I{T1})",'="to reach the requested repeaters"'),
       ("H","J","CONTRACT PROGRESS",f"=IF({NREPS}=0,0,MIN(G{T0}:G{T1})/{NREPS})",f'=MIN(G{T0}:G{T1})&" of "&{NREPS}&" repeaters kitted"')]
ds.row_dimensions[5].height=20; ds.row_dimensions[6].height=54; ds.row_dimensions[7].height=30
for a,b,lab,val,sub in tiles:
    for r in (5,6,7):
        if a!=b: ds.merge_cells(f"{a}{r}:{b}{r}")
        ca=ds[f"{a}{r}"].column; cb=ds[f"{b}{r}"].column
        for c in range(ca,cb+1): ds.cell(r,c).fill=fill(PANEL); ds.cell(r,c).border=Border(left=Side(style="thick",color=ROSE) if c==ca else None)
    ds[f"{a}5"]=lab; ds[f"{a}5"].font=font(10,True,GREY); ds[f"{a}5"].alignment=Alignment(horizontal="left",vertical="bottom",indent=1)
    ds[f"{a}6"]=val; ds[f"{a}6"].font=font(32,True,NAVY); ds[f"{a}6"].alignment=Alignment(horizontal="left",vertical="center",indent=1)
    ds[f"{a}7"]=sub; ds[f"{a}7"].font=font(9,c=GREY); ds[f"{a}7"].alignment=Alignment(horizontal="left",vertical="top",indent=1,wrap_text=True)
ds["B6"].font=font(44,True,ROSE); ds["H6"].number_format="0%"
ds.conditional_formatting.add("D6",FormulaRule(formula=['$D$6="READY"'],font=Font(name=F,size=32,bold=True,color=OK)))
ds.conditional_formatting.add("D6",FormulaRule(formula=['$D$6<>"READY"'],font=Font(name=F,size=32,bold=True,color=WARN)))

# repeater counter row 9
ds["B9"]="REPEATER COUNTER"; ds["B9"].font=font(12,True,NAVY)
ds["D9"]='="Green = kitted · amber = partly released (codes ready) · outlined = requested"'; ds["D9"].font=font(9,c=GREY,i=True)
# counter in a separate strip using columns L.. ? keep within B..J: put on row 10 across E..J? use a dedicated area at right (L)
# Ledger
ds[f"B{T0-1}"]=None
heads=["Sequence","Assy Code","Description","Qty / Rep","Released","Repeaters","Needed for Request","To Release Now","Status"]
for j,h in enumerate(heads): hdr(ds,T0-1,2+j,h)
ds.row_dimensions[T0-1].height=32
for i in range(NP):
    r=T0+i; k=P0+i
    ds[f"B{r}"]=f"=Contract!A{k}"; ds[f"C{r}"]=f"=Contract!B{k}"; ds[f"D{r}"]=f"=Contract!C{k}"; ds[f"E{r}"]=f"=Contract!D{k}"
    ds[f"F{r}"]=f"=Contract!E{k}+SUMIFS({LD},{LB},C{r})"
    ds[f"G{r}"]=f"=IF(E{r}=0,0,INT(F{r}/E{r}))"
    ds[f"H{r}"]=f"={REQ}*E{r}"
    ds[f"I{r}"]=f"=MAX(0,H{r}-F{r})"
    ds[f"J{r}"]=f'=IF(I{r}=0,"READY","RELEASE "&I{r})'
    for col in "BCDEFGHIJ":
        x=ds[f"{col}{r}"]; x.border=BOT; x.font=font(11); x.alignment=L if col in "BCD" else C
    ds[f"C{r}"].font=font(11,True); ds[f"D{r}"].font=font(9,c=GREY); ds[f"G{r}"].font=font(15,True,NAVY); ds[f"I{r}"].font=font(14,True)
    ds.row_dimensions[r].height=23
ds.conditional_formatting.add(f"J{T0}:J{T1}",FormulaRule(formula=[f'$J{T0}="READY"'],fill=fill(OK_S),font=Font(name=F,color=OK,bold=True)))
ds.conditional_formatting.add(f"J{T0}:J{T1}",FormulaRule(formula=[f'$J{T0}<>"READY"'],fill=fill(WARN_S),font=Font(name=F,color=WARN,bold=True)))
ds.conditional_formatting.add(f"I{T0}:I{T1}",CellIsRule(operator="greaterThan",formula=["0"],font=Font(name=F,color=WARN,bold=True,size=14)))
ds.conditional_formatting.add(f"I{T0}:I{T1}",CellIsRule(operator="equal",formula=["0"],font=Font(name=F,color="C3C8D2",size=14)))
ds.conditional_formatting.add(f"G{T0}:G{T1}",FormulaRule(formula=[f"$G{T0}=MIN($G${T0}:$G${T1})"],fill=fill(CRIT_S)))
ds.conditional_formatting.add(f"B{T0}:B{T1}",FormulaRule(formula=[f"$B{T0}=$B{T0-1}"],font=Font(name=F,color="C3C8D2")))
TR=T1+1
ds[f"C{TR}"]="TOTAL"; ds[f"C{TR}"].font=font(11,True)
for col,fx in (("F","SUM"),("G","MIN"),("H","SUM"),("I","SUM")):
    ds[f"{col}{TR}"]=f"={fx}({col}{T0}:{col}{T1})"; ds[f"{col}{TR}"].font=font(12,True,NAVY); ds[f"{col}{TR}"].alignment=C; ds[f"{col}{TR}"].border=Border(top=Side(style="medium",color=NAVY))
ds[f"D{TR}"]="Repeaters column total = lowest code (highlighted) — that code limits the count"; ds[f"D{TR}"].font=font(9,c=GREY,i=True)

# Repeater counter: row 9 is title; boxes in row 10 across columns? Put in column L as vertical strip next to ledger
ds["L9"]=None
ds.column_dimensions["L"].width=20
hdr(ds,T0-1,12,"Repeater")
for n in range(1,NREP+1):
    r=T0+n-1
    if r>T1+6: break
    ds[f"L{r}"]=f'="Rep {n}  ·  "&IF(MIN($G${T0}:$G${T1})>={n},"KITTED",IF(COUNTIF($G${T0}:$G${T1},">="&{n})>0,COUNTIF($G${T0}:$G${T1},">="&{n})&"/{NP} codes","—"))'
    ds[f"L{r}"].alignment=C; ds[f"L{r}"].font=font(10,True)
LC=f"L{T0}:L{T0+NREP-1}"
ds.conditional_formatting.add(LC,FormulaRule(formula=[f'ISNUMBER(SEARCH("KITTED",L{T0}))'],fill=fill(OK),font=Font(name=F,color=WHITE,bold=True)))
ds.conditional_formatting.add(LC,FormulaRule(formula=[f'ISNUMBER(SEARCH("codes",L{T0}))'],fill=fill(WARN_S),font=Font(name=F,color=WARN,bold=True)))
ds.conditional_formatting.add(LC,FormulaRule(formula=[f'ISNUMBER(SEARCH("—",L{T0}))'],fill=fill("ECEEF3"),font=Font(name=F,color=GREY)))
thick=Side(style="medium",color=ROSE)
ds.conditional_formatting.add(LC,FormulaRule(formula=[f'ROW()-{T0-1}<={REQ}'],border=Border(left=thick,right=thick,top=thick,bottom=thick)))
ds["B9"]=None; ds["D9"]=None
ds[f"L{T0-2}"]="REPEATER COUNTER"; ds[f"L{T0-2}"].font=font(12,True,NAVY)

# Latest releases
LR0=TR+3
ds[f"B{LR0-1}"]="LATEST RELEASES"; ds[f"B{LR0-1}"].font=font(12,True,NAVY)
for j,h in enumerate(["Date","Assy Code","Description","Qty","Released By","Logged"]): hdr(ds,LR0,2+j,h)
ds.merge_cells(start_row=LR0,start_column=8,end_row=LR0,end_column=10); ds.cell(LR0,8).value="Note"
hdr(ds,LR0,8,"Note")
CNT=f"MAX(IF(({LB}<>\"\")+({LD}<>\"\"),ROW({LB})))"
ds["L2"]=None
ds["N1"]=None
ds["M1"]=f"=SUMPRODUCT(MAX((({LB}<>\"\")+({LD}<>\"\")>0)*ROW({LB})))"; ds["M1"].font=font(8,c=WHITE)
ds.column_dimensions["M"].hidden=True
for i in range(8):
    r=LR0+1+i; rowexpr=f"$M$1-{i}"
    ok=f'AND({rowexpr}>={R0})'
    for col,src in (("B","A"),("C","B"),("D","C"),("E","D"),("F","E"),("H","F")):
        ds[f"{col}{r}"]=f"=IF({rowexpr}<{R0},\"\",IF(INDEX('Release Log'!${src}:${src},{rowexpr})=\"\",\"\",INDEX('Release Log'!${src}:${src},{rowexpr})))"
        ds[f"{col}{r}"].border=BOT; ds[f"{col}{r}"].font=font(10); ds[f"{col}{r}"].alignment=L if col in "CDH" else C
    ds[f"G{r}"]=f'=IF(B{r}="","","#"&({rowexpr}-{R0-1}))'; ds[f"G{r}"].font=font(9,c=GREY); ds[f"G{r}"].alignment=C; ds[f"G{r}"].border=BOT
    ds.merge_cells(f"H{r}:J{r}")
    ds[f"B{r}"].number_format="dd-mmm-yy"; ds[f"D{r}"].font=font(9,c=GREY); ds[f"E{r}"].font=font(11,True)
ds.freeze_panes="A4"
ds.page_setup.orientation="landscape"; ds.page_setup.fitToWidth=1; ds.page_setup.fitToHeight=1; ds.sheet_properties.pageSetUpPr.fitToPage=True
ds.print_area=f"A1:L{LR0+8}"

# ---------------- How To Use ----------------
hu=wb.create_sheet("How To Use"); hu.sheet_view.showGridLines=False
hu.column_dimensions["B"].width=26; hu.column_dimensions["C"].width=105
hu["B2"]="HOW THE REPEATER COUNTER WORKS"; hu["B2"].font=font(18,True,NAVY)
rows=[("Each time management asks","",True),
 ("1. Set the target","Type the number of repeaters requested in the pink cell on the Dashboard (top right).",False),
 ("2. Release","The 'To Release Now' column tells you how many of each assy code to release to reach the target. Release those work orders and send them to production.",False),
 ("3. Log it","On 'Release Log' add one line per release: date, assy code (drop-down), qty released, your name, optional note (WO no. / repeater). Made a mistake? Add a line with a negative qty.",False),
 ("4. Everyone sees it","The Dashboard updates immediately: Repeaters Kitted, the counter, and Latest Releases. Save the file on SharePoint / Teams so everyone sees the same numbers.",False),
 ("How it counts","",True),
 ("Repeaters per code","Released qty ÷ qty per repeater, rounded down. E.g. erbium spool: 50 released ÷ 24 = 2 repeaters (2 units toward the 3rd).",False),
 ("Repeaters kitted","The lowest value across all assy codes — a repeater only counts when every code has been released for it. The limiting code is highlighted in the table.",False),
 ("Starting point","If some repeaters were released before you started the log, put those quantities in 'Opening Balance' on the Contract sheet.",False),
 ("Changing the contract","Edit codes, qty per repeater or contract repeaters on the Contract sheet (blue cells). Everything follows.",False),
 ("Projecting","Open the Dashboard and press Ctrl+F1 to hide the ribbon. It fits one landscape screen.",False)]
r=4
for k,v,h in rows:
    if h: r+=1; hu[f"B{r}"]=k; hu[f"B{r}"].font=font(12,True,ROSE); r+=1; continue
    hu[f"B{r}"]=k; hu[f"B{r}"].font=font(11,True); hu[f"B{r}"].alignment=Alignment(vertical="top")
    hu[f"C{r}"]=v; hu[f"C{r}"].font=font(11); hu[f"C{r}"].alignment=Alignment(wrap_text=True,vertical="top"); hu.row_dimensions[r].height=34; r+=1
wb.calculation=CalcProperties(fullCalcOnLoad=True)
wb.active=0
wb.save(OUT)

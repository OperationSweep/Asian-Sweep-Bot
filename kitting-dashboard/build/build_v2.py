"""R5B AMP Kitting Control v2 — Microsoft 365, Excel Tables + dynamic arrays, no transaction log.
usage: build_v2.py orig.json out.xlsx [test_mode]
"""
import sys, json, re, datetime, zipfile
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection
from openpyxl.formatting.rule import FormulaRule, CellIsRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo, TableFormula
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.utils import get_column_letter as CL
from openpyxl.workbook.properties import CalcProperties
from openpyxl.comments import Comment

ORIG=json.load(open(sys.argv[1])); OUT=sys.argv[2]
MIGRATION_DATE=datetime.date(2026,10,5)

# ---------- formula helpers (Excel file syntax) ----------
FN={"FILTER":"_xlfn._xlws.FILTER","LET":"_xlfn.LET","HSTACK":"_xlfn.HSTACK","MINIFS":"_xlfn.MINIFS","MAXIFS":"_xlfn.MAXIFS",
    "ISOWEEKNUM":"_xlfn.ISOWEEKNUM","TEXTJOIN":"_xlfn.TEXTJOIN","IFS":"_xlfn.IFS","SEQUENCE":"_xlfn.SEQUENCE","XMATCH":"_xlfn.XMATCH"}
def X(f):
    f=re.sub(r'(?<![\w.])('+"|".join(FN)+r')\(', lambda m: FN[m.group(1)]+"(", f)
    f=re.sub(r'(?<![\w.\[])(v_\w+)', r'_xlpm.\1', f)
    return f
DA_CELLS=[]   # (sheet title, cell) that must be dynamic-array formulas
def da(ws,cell,formula):
    ws[cell]=ArrayFormula(cell,"="+X(formula)); DA_CELLS.append((ws.title,cell))

# ---------- style ----------
NAVY="141A26"; INK="1E2433"; ROSE="B0245E"; ROSE_S="F9DCE7"; GREY="5D6576"; LINE="D9DDE5"; PANEL="F3F4F7"; WHITE="FFFFFF"; MUTE="AEB6C6"
OK="1F8A5B"; OK_S="DDF2E7"; WARN="A8660B"; WARN_S="FBF0DC"; CRIT="C0392B"; CRIT_S="FBE6E3"; INFO="3558B8"; INFO_S="E3EAFA"
INPUT_FILL="FFF2B3"; CALC_FILL="E4E6EB"
F="Arial"
def font(sz=11,b=False,c=INK,i=False): return Font(name=F,size=sz,bold=b,color=c,italic=i)
def fill(c): return PatternFill("solid",start_color=c,end_color=c)
thin=Side(style="thin",color=LINE)
C=Alignment(horizontal="center",vertical="center",wrap_text=True); L=Alignment(horizontal="left",vertical="center",wrap_text=True)
LT=Alignment(horizontal="left",vertical="top",wrap_text=True)
UNLOCK=Protection(locked=False)
def band(ws,cols,rows):
    for r in rows:
        for c in cols: ws.cell(r,c).fill=fill(NAVY)

wb=Workbook()
wb.remove(wb.active)
start=wb.create_sheet("Start Here"); dash=wb.create_sheet("Dashboard"); gantt=wb.create_sheet("Gantt")
cons=wb.create_sheet("Contracts"); asm=wb.create_sheet("Assemblies & Totals"); plan=wb.create_sheet("Weekly Plan")

# ---------- generic table builder ----------
def make_table(ws, name, top, cols, rows, calc, widths, style="TableStyleLight1"):
    """cols: list of (header, kind) kind in {'in','calc'}; calc: {header: formula_text_without_eq (file syntax)}"""
    hr=top; ncol=len(cols)
    # marker row above header
    for j,(h,kind) in enumerate(cols,1):
        m=ws.cell(hr-1,j,"INPUT ▼" if kind=="in" else "CALCULATED – do not type ▼")
        m.font=font(8,True,WARN if kind=="in" else GREY); m.alignment=C
        m.fill=fill(INPUT_FILL if kind=="in" else CALC_FILL)
        hc=ws.cell(hr,j,h); hc.font=font(10,True,INK); hc.alignment=C
        hc.fill=fill(INPUT_FILL if kind=="in" else CALC_FILL)
        hc.border=Border(bottom=Side(style="medium",color=INK))
    for i,row in enumerate(rows):
        r=hr+1+i
        for j,(h,kind) in enumerate(cols,1):
            c=ws.cell(r,j)
            if kind=="calc":
                c.value="="+calc[h]; c.fill=fill(CALC_FILL); c.font=font(10,False,GREY)
            else:
                v=row.get(h); c.value=v; c.font=font(10)
            c.alignment=L if isinstance(c.value,str) and not str(c.value).startswith("=") else C
            c.border=Border(bottom=thin)
    last=hr+max(1,len(rows))
    t=Table(displayName=name, ref=f"A{hr}:{CL(ncol)}{last}")
    t.tableStyleInfo=TableStyleInfo(name=style,showRowStripes=False)
    t._initialise_columns()
    for j,(h,kind) in enumerate(cols):
        t.tableColumns[j].name=h          # must equal the header cell text exactly (Excel requirement)
        if kind=="calc": t.tableColumns[j].calculatedColumnFormula=TableFormula(attr_text=calc[h])
    ws.add_table(t)
    for j,w in enumerate(widths,1): ws.column_dimensions[CL(j)].width=w
    ws.freeze_panes=ws.cell(hr+1,1)
    return hr,last

def T(tbl,col): return f"{tbl}[{col}]"
def TR(tbl,col): return f"{tbl}[[#This Row],[{col}]]"

# =====================================================================
# CONTRACTS
# =====================================================================
cons.sheet_view.showGridLines=False
cons["A1"]="CONTRACTS"; cons["A1"].font=font(18,True,NAVY)
cons["A2"]="One row per contract (top-level ++ code). Yellow columns are typed; grey columns calculate. New rows appear automatically in the Dashboard and Gantt contract lists."
cons["A2"].font=font(10,c=GREY,i=True)
cons["A4"]="SETTING – Update interval (days)"; cons["A4"].font=font(10,True)
cons["C4"]=7; cons["C4"].fill=fill(INPUT_FILL); cons["C4"].font=font(12,True); cons["C4"].alignment=C
cons["D4"]="Supplied totals older than this are flagged as out of date on the Dashboard."; cons["D4"].font=font(9,c=GREY,i=True)
ccols=[("Contract ID","in"),("Friendly name","in"),("Project","in"),("Top-level assembly code","in"),("Amps per repeater","in"),
       ("Total repeaters required","in"),("Kit-release lead (weeks before final build)","in"),("AMPS-build lead (weeks before final build)","in"),
       ("Plan revision","in"),("Plan updated on","in"),("Plan updated by","in"),("Notes","in"),("Setup check","calc")]
ccalc={"Setup check":
 'IF('+TR("tblContracts","Contract ID")+'="","Missing contract ID",'
 'IF(COUNTIF('+T("tblContracts","Contract ID")+','+TR("tblContracts","Contract ID")+')>1,"DUPLICATE contract ID",'
 'IF(OR(NOT(ISNUMBER('+TR("tblContracts","Total repeaters required")+')),'+TR("tblContracts","Total repeaters required")+'<=0),"Total repeaters missing",'
 'IF(OR(NOT(ISNUMBER('+TR("tblContracts","Kit-release lead (weeks before final build)")+')),NOT(ISNUMBER('+TR("tblContracts","AMPS-build lead (weeks before final build)")+'))),"Lead time missing",'
 'IF(COUNTIF(tblAssemblies[Contract],'+TR("tblContracts","Contract ID")+')=0,"Assembly setup required","OK")))))'}
crow=[]
for c in ORIG["contracts"]:
    cid,proj,top,fp,amps,total,note=c[0],c[1],c[2],c[3],c[4],c[5],c[6]
    friendly=f"{proj} · {fp}FP" + ("" if "TBC" not in cid else " (TBC)")
    rev="PLAN export 02-Sep-26" + (" + ASP 01 2026 (Jan-26) weeks W24–W35" if cid=="E2A-R12P-T5-C" else "")
    crow.append({"Contract ID":cid,"Friendly name":friendly,"Project":proj,"Top-level assembly code":top,"Amps per repeater":amps,
                 "Total repeaters required":total,"Kit-release lead (weeks before final build)":c[16] if c[16] is not None else 2,
                 "AMPS-build lead (weeks before final build)":1,"Plan revision":rev,"Plan updated on":MIGRATION_DATE,
                 "Plan updated by":"Migrated from previous workbook","Notes":note})
CH,CL_=make_table(cons,"tblContracts",7,ccols,crow,ccalc,[18,22,24,18,10,12,14,14,30,12,22,26,22])
for r in range(CH+1,CL_+1): cons.cell(r,10).number_format="dd-mmm-yy"
cons.row_dimensions[CH].height=44

# =====================================================================
# ASSEMBLIES & TOTALS
# =====================================================================
asm.sheet_view.showGridLines=False
asm["A1"]="ASSEMBLIES & TOTALS"; asm["A1"].font=font(18,True,NAVY)
asm.merge_cells("A2:L2")
asm["A2"]="Enter the TOTAL supplied to date, not this week's additions. If the existing total is 26 and another 4 are supplied, enter 30."
asm["A2"].font=font(13,True,WHITE); asm["A2"].fill=fill(ROSE); asm["A2"].alignment=L; asm.row_dimensions[2].height=26
asm.merge_cells("A3:L3")
asm["A3"]=("When you change a total, also update 'Updated on', 'Updated by' and 'Source reference' (e.g. the SAP export filename). "
           "If you REDUCE a total to correct a mistake, write why in 'Correction note'. The workbook does not keep previous values.")
asm["A3"].font=font(10,c=INK); asm["A3"].alignment=L; asm.row_dimensions[3].height=30
asm.merge_cells("A4:L4")
asm["A4"]="One row per contract + assembly code. A code used by two contracts needs two rows (separate totals)."
asm["A4"].font=font(9,c=GREY,i=True)
acols=[("Contract","in"),("Assembly group","in"),("Assembly code","in"),("Description","in"),("Qty per repeater","in"),
       ("Cumulative supplied to date","in"),("Updated on","in"),("Updated by","in"),("Source reference","in"),("Correction note","in"),
       ("Complete sets (calc)","calc"),("Check (calc)","calc")]
q=TR("tblAssemblies","Qty per repeater"); s=TR("tblAssemblies","Cumulative supplied to date")
acalc={"Complete sets (calc)":f'IF(AND(ISNUMBER({q}),{q}>0,{q}=INT({q})),INT(IF(ISNUMBER({s}),{s},0)/{q}),"")',
 "Check (calc)":(f'IF({TR("tblAssemblies","Contract")}="","Missing contract",'
   f'IF(COUNTIF(tblContracts[Contract ID],{TR("tblAssemblies","Contract")})=0,"Unknown contract",'
   f'IF({TR("tblAssemblies","Assembly code")}="","Missing assembly code",'
   f'IF(COUNTIFS(tblAssemblies[Contract],{TR("tblAssemblies","Contract")},tblAssemblies[Assembly code],{TR("tblAssemblies","Assembly code")})>1,"DUPLICATE contract + code",'
   f'IF(NOT(AND(ISNUMBER({q}),{q}>0,{q}=INT({q}))),"Qty per repeater missing or invalid",'
   f'IF({s}="","Supplied total missing",IF(OR(NOT(ISNUMBER({s})),{s}<0),"Supplied total invalid",'
   f'IF({TR("tblAssemblies","Updated on")}="","Not yet updated","OK"))))))))')}
# migrate
last_rel={}
for d,cn,code,qty,by,note,desc in ORIG["log"]:
    k=(cn,code); dd=d[:10]; last_rel[k]=max(last_rel.get(k,""),dd)
arows=[]
for cn,grp,code,desc,per,released,_ in ORIG["kits"]:
    lr=last_rel.get((cn,code))
    arows.append({"Contract":cn,"Assembly group":grp,"Assembly code":code,"Description":desc,"Qty per repeater":per,
                  "Cumulative supplied to date":released,
                  "Updated on":datetime.date.fromisoformat(lr) if lr else None,
                  "Updated by":"Migrated from previous workbook" if lr else None,
                  "Source reference":("Previous workbook Release Log (182 lines from COOIS new_e2a, 03-Oct-26)" if lr else "Previous workbook: no releases recorded"),
                  "Correction note":None})
AH,AL=make_table(asm,"tblAssemblies",7,acols,arows,acalc,[18,14,17,40,10,13,12,18,34,22,11,24])
for r in range(AH+1,AL+1): asm.cell(r,7).number_format="dd-mmm-yy"
asm.row_dimensions[AH].height=34

# =====================================================================
# WEEKLY PLAN
# =====================================================================
plan.sheet_view.showGridLines=False
plan["A1"]="WEEKLY PLAN — FINAL BUILD / HOUSING WEEKS"; plan["A1"].font=font(18,True,NAVY)
plan.merge_cells("A2:F2")
plan["A2"]=("Each row = repeaters planned for FINAL BUILD (housing) in the week starting on that Monday. "
            "Kit release and AMPS build weeks are calculated from the lead times on 'Contracts'.")
plan["A2"].font=font(11,True,WHITE); plan["A2"].fill=fill(INFO); plan["A2"].alignment=L; plan.row_dimensions[2].height=30
plan.merge_cells("A3:F3")
plan["A3"]=("Move: change the date · Split: change the qty and add a row · Cancel: delete the row (or set qty 0) · "
            "Replace a whole schedule: filter Contract, delete ALL its rows, then paste the new rows. Never leave old and new rows together. "
            "Record the new plan revision on 'Contracts'.")
plan["A3"].font=font(9,c=INK); plan["A3"].alignment=L; plan.row_dimensions[3].height=30
pcols=[("Contract","in"),("Final build week (Monday)","in"),("Repeaters planned","in"),("Note","in"),("ISO week (calc)","calc"),("Check (calc)","calc")]
d=TR("tblPlan","Final build week (Monday)"); pq=TR("tblPlan","Repeaters planned")
pcalc={"ISO week (calc)":f'IF(ISNUMBER({d}),"W"&TEXT(ISOWEEKNUM({d}),"00")&"-"&YEAR({d}-WEEKDAY({d},2)+4),"")'.replace("ISOWEEKNUM","_xlfn.ISOWEEKNUM"),
 "Check (calc)":(f'IF({TR("tblPlan","Contract")}="","Missing contract",'
   f'IF(COUNTIF(tblContracts[Contract ID],{TR("tblPlan","Contract")})=0,"Unknown contract",'
   f'IF(NOT(ISNUMBER({d})),"Missing date",IF(WEEKDAY({d},2)<>1,"Not a Monday",'
   f'IF(OR(NOT(ISNUMBER({pq})),{pq}<0),"Qty missing or invalid","OK")))))')}
prows=[]
for cn,yy,ww,qq,ws_,note in ORIG["plan"]:
    dt=datetime.date.fromisoformat(ws_[:10])
    prows.append({"Contract":cn,"Final build week (Monday)":dt,"Repeaters planned":qq,"Note":note})
PH,PL=make_table(plan,"tblPlan",6,pcols,prows,pcalc,[22,16,12,46,12,20])
for r in range(PH+1,PL+1): plan.cell(r,2).number_format="ddd dd-mmm-yy"
plan.row_dimensions[PH].height=30

# =====================================================================
# Names
# =====================================================================
def name(n,ref): wb.defined_names[n]=DefinedName(n,attr_text=ref)
name("ContractList","tblContracts[Contract ID]")
name("ThisWeek","TODAY()-WEEKDAY(TODAY(),2)+1")
name("UpdateIntervalDays","Contracts!$C$4")

# =====================================================================
# Data validation (applied to whole table columns + room to grow)
# =====================================================================
GROW=5000
def dv(ws,rule,rng,err,title="Check entry",style="stop"):
    rule.error=err; rule.errorTitle=title; rule.showErrorMessage=True; rule.errorStyle=style
    ws.add_data_validation(rule); rule.add(rng)
dv(asm,DataValidation(type="list",formula1="ContractList",allow_blank=True),f"A{AH+1}:A{AH+GROW}","Pick a contract that exists on the Contracts sheet.")
dv(plan,DataValidation(type="list",formula1="ContractList",allow_blank=True),f"A{PH+1}:A{PH+GROW}","Pick a contract that exists on the Contracts sheet.")
dv(asm,DataValidation(type="whole",operator="greaterThan",formula1="0",allow_blank=True),f"E{AH+1}:E{AH+GROW}","Qty per repeater must be a whole number above 0.")
dv(asm,DataValidation(type="whole",operator="greaterThanOrEqual",formula1="0",allow_blank=True),f"F{AH+1}:F{AH+GROW}","Enter the cumulative total supplied to date (whole number, 0 or more).")
dv(asm,DataValidation(type="date",operator="greaterThan",formula1="43831",allow_blank=True),f"G{AH+1}:G{AH+GROW}","Enter a date, e.g. 05/10/2026.")
dv(plan,DataValidation(type="custom",formula1=f"AND(ISNUMBER(B{PH+1}),WEEKDAY(B{PH+1},2)=1)",allow_blank=True),f"B{PH+1}:B{PH+GROW}","Enter the MONDAY date of the final build week.","Monday date required")
dv(plan,DataValidation(type="whole",operator="greaterThanOrEqual",formula1="0",allow_blank=True),f"C{PH+1}:C{PH+GROW}","Repeaters planned must be a whole number (0 or more).")
for col in "EF": dv(cons,DataValidation(type="whole",operator="greaterThan",formula1="0",allow_blank=True),f"{col}{CH+1}:{col}{CH+GROW}","Whole number above 0.")
for col in "GH": dv(cons,DataValidation(type="whole",operator="between",formula1="0",formula2="20",allow_blank=True),f"{col}{CH+1}:{col}{CH+GROW}","Lead time in whole weeks (0–20).")
dv(cons,DataValidation(type="date",operator="greaterThan",formula1="43831",allow_blank=True),f"J{CH+1}:J{CH+GROW}","Enter a date.")
dv(cons,DataValidation(type="whole",operator="between",formula1="1",formula2="365"),"C4","Days between 1 and 365.")
# check-column CF (whole growth range)
for ws,colL,top in ((cons,"M",CH),(asm,"L",AH),(plan,"F",PH)):
    rng=f"{colL}{top+1}:{colL}{top+GROW}"
    ws.conditional_formatting.add(rng,FormulaRule(formula=[f'AND({colL}{top+1}<>"",{colL}{top+1}<>"OK",{colL}{top+1}<>"Not yet updated")'],fill=fill(CRIT_S),font=Font(name=F,bold=True,color=CRIT)))
    ws.conditional_formatting.add(rng,FormulaRule(formula=[f'{colL}{top+1}="Not yet updated"'],fill=fill(WARN_S),font=Font(name=F,bold=True,color=WARN)))
    ws.conditional_formatting.add(rng,FormulaRule(formula=[f'{colL}{top+1}="OK"'],font=Font(name=F,color=OK)))

# =====================================================================
# DASHBOARD
# =====================================================================
ds=dash; ds.sheet_view.showGridLines=False; ds.sheet_view.zoomScale=85
widths=[2,17,17,36,10,12,11,13,13,13,12,26,2]
for j,w in enumerate(widths,1): ds.column_dimensions[CL(j)].width=w
band(ds,range(2,13),(1,2,3))
ds.row_dimensions[2].height=34; ds.row_dimensions[3].height=28
ds["B2"]="R5B AMP KITTING CONTROL"; ds["B2"].font=font(22,True,WHITE); ds["B2"].alignment=Alignment(vertical="center")
ds["B3"]="CONTRACT ▸"; ds["B3"].font=font(10,True,MUTE); ds["B3"].alignment=Alignment(horizontal="right",vertical="center")
ds.merge_cells("C3:D3"); SEL="$C$3"
ds["C3"]="E2A-R12P-T5-C"; ds["C3"].font=font(15,True,NAVY); ds["C3"].fill=fill(INPUT_FILL); ds["C3"].alignment=C; ds["C3"].protection=UNLOCK
ds["D3"].fill=fill(INPUT_FILL)
dv(ds,DataValidation(type="list",formula1="ContractList",allow_blank=False),"C3","Pick a contract from the list (Contracts sheet).")
CI=f'MATCH({SEL},tblContracts[Contract ID],0)'
def cv(col): return f'INDEX(tblContracts[{col}],{CI})'
ds.merge_cells("E3:I3")
ds["E3"]=f'=IFERROR({cv("Friendly name")}&"  ·  "&{cv("Top-level assembly code")}&"  ·  "&{cv("Amps per repeater")}&" amps per repeater","Contract not found — pick from the list")'
ds["E3"].font=font(11,True,WHITE); ds["E3"].alignment=Alignment(vertical="center")
ds.merge_cells("J3:L3")
ds["J3"]='="Today "&TEXT(TODAY(),"dd-mmm-yy")&"  ·  week W"&TEXT(_xlfn.ISOWEEKNUM(TODAY()),"00")&" (starts "&TEXT(ThisWeek,"dd-mmm")&")"'
ds["J3"].font=font(10,c=MUTE); ds["J3"].alignment=Alignment(horizontal="right",vertical="center")
# definition banner
ds.merge_cells("B5:L5")
ds["B5"]="One complete repeater kit means every required assembly has been supplied. It does not mean the repeater has finished production."
ds["B5"].font=font(11,True,INFO); ds["B5"].fill=fill(INFO_S); ds["B5"].alignment=L; ds.row_dimensions[5].height=24

# core values (visible, row 7-9 tiles). Keep calculations in cells referenced by others.
A_C="tblAssemblies[Contract]"; A_S="tblAssemblies[Complete sets (calc)]"
n_asm=f'COUNTIF({A_C},{SEL})'
n_bad=f'COUNTIFS({A_C},{SEL},{A_S},"")'
KITS=f'IF({n_asm}=0,"Assembly setup required",IF(OR({n_bad}>0,P10>0),"Setup incomplete",MINIFS({A_S},{A_C},{SEL})))'
TOTAL=f'IFERROR({cv("Total repeaters required")},0)'
KL=f'IFERROR({cv("Kit-release lead (weeks before final build)")},0)'
AL=f'IFERROR({cv("AMPS-build lead (weeks before final build)")},0)'
P_C="tblPlan[Contract]"; P_D="tblPlan[Final build week (Monday)]"; P_Q="tblPlan[Repeaters planned]"
# hidden calc cells (column O..)  -> keep on sheet, hidden columns
ds.column_dimensions["O"].hidden=True; ds.column_dimensions["P"].hidden=True
calc_cells={
 "P1":("Kits (number or message)",KITS),
 "P2":("Total required",TOTAL),
 "P3":("Kit lead (weeks)",KL),
 "P4":("AMPS lead (weeks)",AL),
 "P5":("Kits due by this week",f'SUMIFS({P_Q},{P_C},{SEL},{P_D},"<"&(ThisWeek+7*(P3+1)))'),
 "P6":("Kits due next 4 weeks",f'SUMIFS({P_Q},{P_C},{SEL},{P_D},">="&(ThisWeek+7*(P3+1)),{P_D},"<"&(ThisWeek+7*(P3+5)))'),
 "P7":("Scheduled total",f'SUMIFS({P_Q},{P_C},{SEL})'),
 "P8":("Target in use",f'IF(ISNUMBER($K$18),$K$18,P5+P6)'),
 "P9":("Assemblies count",n_asm),
 "P10":("Setup issues (qty/dup/unknown)",f'COUNTIFS({A_C},{SEL},tblAssemblies[Check (calc)],"<>OK",tblAssemblies[Check (calc)],"<>Not yet updated",tblAssemblies[Check (calc)],"<>Supplied total missing")'),
 "P11":("Assemblies without update date",f'COUNTIFS({A_C},{SEL},tblAssemblies[Updated on],"")'),
 "P12":("Latest update",f'IF(P9=0,"",MAXIFS(tblAssemblies[Updated on],{A_C},{SEL}))'),
 "P13":("Oldest update",f'IF(P9-P11<=0,"",MINIFS(tblAssemblies[Updated on],{A_C},{SEL},tblAssemblies[Updated on],"<>"))'),
 "P14":("Supplied totals missing",f'COUNTIFS({A_C},{SEL},tblAssemblies[Cumulative supplied to date],"")'),
}
for cell,(lab,f_) in calc_cells.items():
    r=int(cell[1:]); ds[f"O{r}"]=lab; ds[cell]="="+X(f_)
# shortage cells (dynamic arrays)
da(ds,"P15",f'IF(P9=0,0,LET(v_f,{A_C}={SEL},v_q,FILTER(tblAssemblies[Qty per repeater],v_f),v_s,FILTER(tblAssemblies[Cumulative supplied to date],v_f),'
    'v_ok,ISNUMBER(v_q)*(v_q>0),v_s0,IF(ISNUMBER(v_s),v_s,0),v_g,IF(v_ok,P8*v_q-v_s0,0),SUM(IF(v_g>0,v_g,0))))'); ds["O15"]="Additional assemblies needed"
da(ds,"P16",f'IF(P9=0,0,LET(v_f,{A_C}={SEL},v_q,FILTER(tblAssemblies[Qty per repeater],v_f),v_s,FILTER(tblAssemblies[Cumulative supplied to date],v_f),'
    'v_s0,IF(ISNUMBER(v_s),v_s,0),v_g,IF(ISNUMBER(v_q)*(v_q>0),P8*v_q-v_s0,0),SUM(--(v_g>0))))'); ds["O16"]="Codes short"
da(ds,"P17",f'IF(P15=0,"",LET(v_f,{A_C}={SEL},v_c,FILTER(tblAssemblies[Assembly code],v_f),v_q,FILTER(tblAssemblies[Qty per repeater],v_f),v_s,FILTER(tblAssemblies[Cumulative supplied to date],v_f),'
    'v_s0,IF(ISNUMBER(v_s),v_s,0),v_g,IF(ISNUMBER(v_q)*(v_q>0),P8*v_q-v_s0,0),v_m,MAX(v_g),INDEX(v_c,MATCH(v_m,v_g,0))&" (+"&v_m&")"))'); ds["O17"]="Largest gap"
ds["P18"]="=IF(P13=\"\",\"\",TODAY()-P13)"; ds["O18"]="Age of oldest update (days)"; ds["P18"].number_format="0"
ds["O19"]="Next kit-release week (Monday)"; ds["P19"]="="+X(f'IFERROR(1/(1/MINIFS({P_D},{P_C},{SEL},{P_D},">="&(ThisWeek+7*(P3+1)))),"")')
ds["O20"]="Repeaters in that week"; ds["P20"]="="+X(f'IF(P19="","",SUMIFS({P_Q},{P_C},{SEL},{P_D},">="&P19,{P_D},"<"&(P19+7)))')
for _r in (19,20): ds[f"O{_r}"].font=font(8,c=GREY); ds[f"P{_r}"].font=font(8,c=GREY)
for r in range(1,19):
    ds[f"O{r}"].font=font(8,c=GREY); ds[f"P{r}"].font=font(8,c=GREY)

# Tiles rows 7-9
tiles=[("B","C","COMPLETE REPEATER KITS RELEASED",'=IF(ISNUMBER(P1),P1&" of "&P2,P1)','=IF(ISNUMBER(P1),TEXT(IF(P2>0,P1/P2,0),"0%")&" of the "&P2&" required","See setup message below")'),
       ("D","D","KITS DUE BY THIS WEEK","=P5",'="Compared with current plan · final build up to W"&TEXT(_xlfn.ISOWEEKNUM(ThisWeek+7*P3),"00")'),
       ("E","G","AHEAD / BEHIND CURRENT PLAN",'=IF(ISNUMBER(P1),IF(P1-P5>0,"+"&(P1-P5),IF(P1-P5=0,"0",""&(P1-P5))),"–")',
            '=IF(ISNUMBER(P1),IF(P1-P5>0,"Ahead of current plan",IF(P1-P5=0,"On current plan","Behind current plan")),"Not available until setup is complete")'),
       ("H","I","KITS DUE OVER THE NEXT FOUR WEEKS","=P6",'="Kit-release weeks W"&TEXT(_xlfn.ISOWEEKNUM(ThisWeek+7),"00")&"–W"&TEXT(_xlfn.ISOWEEKNUM(ThisWeek+28),"00")'),
       ("J","L","ADDITIONAL ASSEMBLIES NEEDED TO REACH TARGET","=P15",'="Units across "&P16&" assembly codes · target "&P8&" complete kits"')]
ds.row_dimensions[7].height=30; ds.row_dimensions[8].height=46; ds.row_dimensions[9].height=30
for a,b,lab,val,sub in tiles:
    for r in (7,8,9):
        if a!=b: ds.merge_cells(f"{a}{r}:{b}{r}")
        ca=ds[f"{a}{r}"].column; cb=ds[f"{b}{r}"].column
        for c in range(ca,cb+1):
            ds.cell(r,c).fill=fill(PANEL); ds.cell(r,c).border=Border(left=Side(style="thick",color=ROSE) if c==ca else None)
    ds[f"{a}7"]=lab; ds[f"{a}7"].font=font(9,True,GREY); ds[f"{a}7"].alignment=Alignment(horizontal="left",vertical="bottom",indent=1,wrap_text=True)
    ds[f"{a}8"]=val; ds[f"{a}8"].font=font(26,True,NAVY); ds[f"{a}8"].alignment=Alignment(horizontal="left",vertical="center",indent=1)
    ds[f"{a}9"]=sub; ds[f"{a}9"].font=font(9,c=GREY); ds[f"{a}9"].alignment=Alignment(horizontal="left",vertical="top",indent=1,wrap_text=True)
ds["B8"].font=font(26,True,ROSE)
ds.conditional_formatting.add("E8",FormulaRule(formula=['LEFT($E$8,1)="+"'],font=Font(name=F,size=26,bold=True,color=OK)))
ds.conditional_formatting.add("E8",FormulaRule(formula=['LEFT($E$8,1)="-"'],font=Font(name=F,size=26,bold=True,color=CRIT)))
ds.conditional_formatting.add("B8",FormulaRule(formula=['NOT(ISNUMBER($P$1))'],font=Font(name=F,size=14,bold=True,color=CRIT)))

# progress + messages rows 11-16
def msg(r,label,formula,size=10,bold=False):
    ds[f"B{r}"]=label; ds[f"B{r}"].font=font(9,True,GREY); ds[f"B{r}"].alignment=Alignment(vertical="center")
    ds.merge_cells(f"C{r}:L{r}"); ds[f"C{r}"]=formula; ds[f"C{r}"].font=font(size,bold); ds[f"C{r}"].alignment=L
    ds.row_dimensions[r].height=20
msg(11,"PROGRESS",'=IF(ISNUMBER(P1),REPT("█",ROUND(30*MIN(1,IF(P2>0,P1/P2,0)),0))&REPT("░",30-ROUND(30*MIN(1,IF(P2>0,P1/P2,0)),0))&"   "&P1&" of "&P2&" complete repeater kits","–")',11,True)
ds["C11"].font=Font(name="Consolas",size=11,bold=True,color=ROSE)
msg(12,"SCHEDULE CHECK",'=IF(P7=P2,"Scheduled total matches the contract ("&P7&" repeaters)",IF(P7<P2,(P2-P7)&" repeaters still need scheduling ("&P7&" scheduled of "&P2&")","Plan exceeds contract total by "&(P7-P2)&" ("&P7&" scheduled vs "&P2&" required)"))',10,True)
msg(13,"PLAN REVISION",f'=IFERROR(IF({cv("Plan revision")}&""="","No plan revision recorded on Contracts",{cv("Plan revision")})&"  ·  updated "&IF(ISNUMBER({cv("Plan updated on")}),TEXT({cv("Plan updated on")},"dd-mmm-yy"),"(date missing)")&" by "&IF({cv("Plan updated by")}&""="","(name missing)",{cv("Plan updated by")}),"")')
msg(14,"SUPPLIED TOTALS",'=IF(P9=0,"No assemblies set up",IF(P11=P9,"No update dates recorded for any assembly (incomplete information)",'
    '"Last updated "&TEXT(P12,"dd-mmm-yy")&IF(P13<>P12,"  ·  oldest "&TEXT(P13,"dd-mmm-yy")&" (mixed dates)","  ·  all assemblies same date")'
    '&IF(P11>0,"  ·  "&P11&" assemblies have no update date","")&IF(P18>UpdateIntervalDays,"  ·  OUT OF DATE (older than "&UpdateIntervalDays&" days)","")))')
msg(15,"SETUP",'=IF(P9=0,"Assembly setup required — add this contract\'s assemblies on \'Assemblies & Totals\'",IF(P10>0,P10&" setup problem(s) — see Check column on \'Assemblies & Totals\'",IF(P14>0,P14&" assemblies have no supplied total","Setup OK ("&P9&" required assemblies)")))')
msg(16,"NEXT ACTION",'=IF(P9=0,"Add the required assemblies for "&$C$3&" on \'Assemblies & Totals\' (contract, group, code, description, qty per repeater).",'
    'IF(P10>0,"Fix the "&P10&" setup problem(s) on \'Assemblies & Totals\' — filter the Check column for anything other than OK.",'
    'IF(OR(P11>0,P14>0,AND(ISNUMBER(P18),P18>UpdateIntervalDays)),"Update cumulative supplied totals on \'Assemblies & Totals\' from the latest verified source, then fill Updated on / by / Source reference.",'
    'IF(P15>0,"To reach the target of "&P8&" complete kits, supply "&P15&" more assemblies across "&P16&" codes. Largest gap: "&P17&".",'
    'IF(P8=0,"No kits are due by this week or in the next four weeks (current plan)."&IF(P19="","  No later kit-release weeks are scheduled."," Next scheduled kit release: week W"&TEXT(_xlfn.ISOWEEKNUM(P19-7*P3),"00")&" ("&P20&" repeaters, final build W"&TEXT(_xlfn.ISOWEEKNUM(P19),"00")&")."),'
    '"Supplied totals cover the target of "&P8&" complete kits. No additional assemblies needed for this target.")))))',11,True)
ds["C16"].fill=fill(WARN_S); ds["B16"].font=font(9,True,WARN)
for r in (12,):
    ds.conditional_formatting.add(f"C{r}",FormulaRule(formula=['$P$7<>$P$2'],font=Font(name=F,bold=True,color=CRIT)))
    ds.conditional_formatting.add(f"C{r}",FormulaRule(formula=['$P$7=$P$2'],font=Font(name=F,bold=True,color=OK)))
ds.conditional_formatting.add("C14",FormulaRule(formula=['ISNUMBER(SEARCH("OUT OF DATE",$C$14))+ISNUMBER(SEARCH("no update",$C$14))'],font=Font(name=F,bold=True,color=WARN)))
ds.conditional_formatting.add("C15",FormulaRule(formula=['LEFT($C$15,8)<>"Setup OK"'],font=Font(name=F,bold=True,color=CRIT)))

# target row 18
ds["B18"]="TARGET"; ds["B18"].font=font(9,True,GREY)
ds.merge_cells("C18:J18")
ds["C18"]="Target complete repeater kits (type a number, or leave blank to use kits due by this week + next 4 weeks) ▸"
ds["C18"].font=font(10,True,INK); ds["C18"].alignment=Alignment(horizontal="right",vertical="center")
ds["K18"]=None; ds["K18"].fill=fill(INPUT_FILL); ds["K18"].font=font(14,True); ds["K18"].alignment=C; ds["K18"].protection=UNLOCK
ds["K18"].border=Border(left=Side(style="medium",color=WARN),right=Side(style="medium",color=WARN),top=Side(style="medium",color=WARN),bottom=Side(style="medium",color=WARN))
dv(ds,DataValidation(type="whole",operator="greaterThanOrEqual",formula1="0",allow_blank=True),"K18","Target must be a whole number of repeater kits.")
ds["L18"]='="In use: "&P8&IF(ISNUMBER($K$18)," (entered)"," (from plan)")'; ds["L18"].font=font(10,True,WARN)
ds.row_dimensions[18].height=26

# assembly list header row 20, spill row 21
LH=20
heads=["Assembly group","Assembly code","Description","Qty per repeater","Supplied to date","Complete sets","Toward next set","Needed for target","Additional needed","Updated on","Check"]
for j,h in enumerate(heads):
    c=ds.cell(LH,2+j,h); c.font=font(9,True,WHITE); c.fill=fill(INK); c.alignment=C
ds.row_dimensions[LH].height=30
ds["B19"]="REQUIRED ASSEMBLIES (actual cumulative supplied totals)"; ds["B19"].font=font(12,True,NAVY)
ds["H19"]="Needed / Additional use the target above"; ds["H19"].font=font(8,c=GREY,i=True)
LIST=(f'IF(P9=0,"Assembly setup required",LET(v_f,{A_C}={SEL},'
      'v_g,FILTER(tblAssemblies[Assembly group],v_f),v_c,FILTER(tblAssemblies[Assembly code],v_f),v_d,FILTER(tblAssemblies[Description],v_f),'
      'v_q,FILTER(tblAssemblies[Qty per repeater],v_f),v_s,FILTER(tblAssemblies[Cumulative supplied to date],v_f),'
      'v_u,FILTER(tblAssemblies[Updated on],v_f),v_k,FILTER(tblAssemblies[Check (calc)],v_f),'
      'v_ok,ISNUMBER(v_q)*(v_q>0),v_s0,IF(ISNUMBER(v_s),v_s,0),'
      'v_sets,IF(v_ok,INT(v_s0/v_q),"check qty"),v_tn,IF(v_ok,MOD(v_s0,v_q)&" of "&v_q,""),'
      'v_need,IF(v_ok,P8*v_q,""),v_add,IF(v_ok,IF(P8*v_q-v_s0>0,P8*v_q-v_s0,0),""),'
      'HSTACK(v_g,v_c,v_d,IF(ISNUMBER(v_q),v_q,"missing"),IF(ISNUMBER(v_s),v_s,"missing"),v_sets,v_tn,v_need,v_add,IF(ISNUMBER(v_u),v_u,"not set"),v_k)))')
da(ds,"B21",LIST)
SP=21; SPE=SP+3000
for col,fmt in (("E","0"),("F","#,##0"),("G","0"),("I","#,##0"),("J","#,##0"),("K","dd-mmm-yy")):
    for r in range(SP,SP+400): ds[f"{col}{r}"].number_format=fmt
for r in range(SP,SP+400):
    for c in range(2,13): ds.cell(r,c).font=font(10); ds.cell(r,c).alignment=C if c not in (2,3,4,12) else L
rng=f"B{SP}:L{SPE}"
ds.conditional_formatting.add(rng,FormulaRule(formula=[f'$C{SP}<>""'],border=Border(bottom=thin)))
ds.conditional_formatting.add(f"G{SP}:G{SPE}",FormulaRule(formula=[f'AND(ISNUMBER($G{SP}),ISNUMBER($P$1),$G{SP}=$P$1)'],fill=fill(CRIT_S),font=Font(name=F,bold=True,color=CRIT)))
ds.conditional_formatting.add(f"J{SP}:J{SPE}",FormulaRule(formula=[f'AND(ISNUMBER($J{SP}),$J{SP}>0)'],fill=fill(WARN_S),font=Font(name=F,bold=True,color=WARN)))
ds.conditional_formatting.add(f"J{SP}:J{SPE}",FormulaRule(formula=[f'AND(ISNUMBER($J{SP}),$J{SP}=0)'],font=Font(name=F,bold=True,color=OK)))
ds.conditional_formatting.add(f"L{SP}:L{SPE}",FormulaRule(formula=[f'AND($L{SP}<>"",$L{SP}<>"OK")'],font=Font(name=F,bold=True,color=CRIT)))
ds.conditional_formatting.add(f"B{SP}:B{SPE}",FormulaRule(formula=[f'AND($B{SP}<>"",$B{SP}=$B{SP-1})'],font=Font(name=F,color="C3C8D2")))
ds["M19"]=None
ds.freeze_panes="A5"
ds.protection.sheet=True; ds.protection.formatColumns=False; ds.protection.formatRows=False; ds.protection.autoFilter=False
ds.page_setup.orientation="landscape"; ds.page_setup.fitToWidth=1; ds.page_setup.fitToHeight=0; ds.sheet_properties.pageSetUpPr.fitToPage=True

# =====================================================================
# GANTT (selected contract, 52-week window, plain SUMIFS)
# =====================================================================
gt=gantt; gt.sheet_view.showGridLines=False; gt.sheet_view.zoomScale=85
GW=52; G0=5
gt.column_dimensions["A"].width=2; gt.column_dimensions["B"].width=30; gt.column_dimensions["C"].width=18; gt.column_dimensions["D"].width=9
for k in range(GW): gt.column_dimensions[CL(G0+k)].width=6.2
GL=CL(G0+GW-1)
band(gt,range(2,G0+GW),(1,2,3))
gt.row_dimensions[2].height=30
gt["B2"]="R5B AMP KITTING — WEEKLY SCHEDULE"; gt["B2"].font=font(18,True,WHITE)
gt["B3"]="CONTRACT ▸"; gt["B3"].font=font(10,True,MUTE); gt["B3"].alignment=Alignment(horizontal="right",vertical="center")
gt["C3"]="E2A-R12P-T5-C"; gt["C3"].font=font(12,True,NAVY); gt["C3"].fill=fill(INPUT_FILL); gt["C3"].alignment=C; gt["C3"].protection=UNLOCK
gt.merge_cells("C3:D3"); gt["D3"].fill=fill(INPUT_FILL)
dv(gt,DataValidation(type="list",formula1="ContractList",allow_blank=False),"C3","Pick a contract from the list.")
GS="$C$3"; GCI=f'MATCH({GS},tblContracts[Contract ID],0)'
def gcv(col): return f'INDEX(tblContracts[{col}],{GCI})'
gt.merge_cells("E3:Z3")
gt["E3"]=f'=IFERROR({gcv("Friendly name")}&"  ·  kit release "&{gcv("Kit-release lead (weeks before final build)")}&" wk before final build  ·  AMPS build "&{gcv("AMPS-build lead (weeks before final build)")}&" wk before","Contract not found")'
gt["E3"].font=font(10,True,WHITE)
# window controls row 5
gt["B5"]="Weeks before this week ▸"; gt["B5"].font=font(9,True,GREY); gt["B5"].alignment=Alignment(horizontal="right")
gt["C5"]=6; gt["C5"].fill=fill(INPUT_FILL); gt["C5"].protection=UNLOCK; gt["C5"].alignment=C; gt["C5"].font=font(11,True)
dv(gt,DataValidation(type="whole",operator="between",formula1="0",formula2="520",allow_blank=False),"C5","Whole number of weeks (0–520).")
gt["B6"]="…or fixed start date ▸"; gt["B6"].font=font(9,True,GREY); gt["B6"].alignment=Alignment(horizontal="right")
gt["C6"]=None; gt["C6"].fill=fill(INPUT_FILL); gt["C6"].protection=UNLOCK; gt["C6"].number_format="dd-mmm-yy"; gt["C6"].alignment=C
dv(gt,DataValidation(type="date",operator="greaterThan",formula1="43831",allow_blank=True),"C6","Enter a date, or leave blank.")
gt["D5"]='=IF(ISNUMBER($C$6),$C$6-WEEKDAY($C$6,2)+1,ThisWeek-7*$C$5)'; gt["D5"].number_format="dd-mmm-yy"; gt["D5"].font=font(8,c=GREY)
gt["D6"]="window start"; gt["D6"].font=font(7,c=GREY)
WS="$D$5"
KLg=f'IFERROR({gcv("Kit-release lead (weeks before final build)")},0)'; ALg=f'IFERROR({gcv("AMPS-build lead (weeks before final build)")},0)'
gt.merge_cells("E5:"+GL+"5")
gt["E5"]=(f'="Final builds scheduled outside this window:  before = "&SUMIFS({P_Q},{P_C},{GS},{P_D},"<"&{WS})&'
          f'"   ·   after = "&SUMIFS({P_Q},{P_C},{GS},{P_D},">="&({WS}+7*{GW}))&'
          f'"   ·   total scheduled = "&SUMIFS({P_Q},{P_C},{GS})&" of "&IFERROR({gcv("Total repeaters required")},0)&" required (all dates are included in Dashboard calculations)"')
gt["E5"].font=font(9,True,INFO)
gt.merge_cells("E6:"+GL+"6")
gt["E6"]=("Kit release, AMPS build and final build show SCHEDULED quantities from the Weekly Plan. Supplied totals show how many kits are covered, "
          "not which weekly batch was released, and say nothing about AMPS-build or final-build completion.")
gt["E6"].font=font(8,c=GREY,i=True); gt["E6"].alignment=L; gt.row_dimensions[6].height=24
# legend row 7
leg=[("E7","Kit release (scheduled)",ROSE,ROSE_S),("I7","Kit covered by supplied totals",OK,OK_S),("N7","Kit due, NOT covered",WHITE,CRIT),
     ("R7","AMPS build (scheduled)",WARN,WARN_S),("V7","Final build / housing (scheduled)",INFO,INFO_S),("AA7","│ red line = this week",CRIT,"FFFFFF")]
for cell,t,fc,bc in leg:
    x=gt[cell]; x.value=t; x.font=font(8,True,fc); x.fill=fill(bc); x.alignment=C
    c0=x.column; gt.merge_cells(start_row=7,start_column=c0,end_row=7,end_column=c0+3)
# header rows 9 (ISO year), 10 (week), 11 (Monday)
HY,HW,HD=9,10,11
for r,lab in ((HY,"ISO year"),(HW,"ISO week"),(HD,"Week starts (Mon)")):
    gt[f"C{r}"]=lab; gt[f"C{r}"].font=font(8,True,MUTE); gt[f"C{r}"].alignment=Alignment(horizontal="right")
    for c in range(2,5): gt.cell(r,c).fill=fill(INK)
for k in range(GW):
    col=CL(G0+k)
    gt[f"{col}{HD}"]=f"={WS}+{7*k}"; gt[f"{col}{HD}"].number_format="dd-mmm"
    gt[f"{col}{HW}"]=f'="W"&TEXT(_xlfn.ISOWEEKNUM({col}{HD}),"00")'
    gt[f"{col}{HY}"]=f'=IF(OR({k}=0,_xlfn.ISOWEEKNUM({col}{HD})=1),YEAR({col}{HD}+3),"")'
    for r,fz,fc,b in ((HY,8,MUTE,True),(HW,9,WHITE,True),(HD,7,MUTE,False)):
        x=gt[f"{col}{r}"]; x.font=font(fz,b,fc); x.fill=fill(INK); x.alignment=C
# stage rows
R_K,R_A,R_H,R_CUM,R_STAT=13,14,15,17,18
stages=[(R_K,"Kit release","W − kit lead",ROSE,KLg),(R_A,"AMPS build","W − AMPS lead",WARN,ALg),(R_H,"Final build / housing","W (plan week)",INFO,"0")]
for r,lab,sub,colr,lead in stages:
    gt[f"B{r}"]=lab; gt[f"B{r}"].font=font(10,True,colr); gt[f"C{r}"]=sub; gt[f"C{r}"].font=font(8,c=GREY)
    gt[f"D{r}"]=f'=SUMIFS({P_Q},{P_C},{GS},{P_D},">="&({WS}+7*({lead})),{P_D},"<"&({WS}+7*({lead}+{GW})))'
    gt[f"D{r}"].font=font(9,True); gt[f"D{r}"].alignment=C
    for k in range(GW):
        col=CL(G0+k)
        sx=f'SUMIFS({P_Q},{P_C},{GS},{P_D},">="&({col}${HD}+7*({lead})),{P_D},"<"&({col}${HD}+7*({lead})+7))'
        gt[f"{col}{r}"]=f'=IF({sx}=0,"",{sx})'
        gt[f"{col}{r}"].alignment=C; gt[f"{col}{r}"].font=font(9,True)
    gt.row_dimensions[r].height=20
gt["D12"]="in view"; gt["D12"].font=font(7,c=GREY); gt["D12"].alignment=C
# cumulative kits due vs covered
gt[f"B{R_CUM}"]="Cumulative kits due"; gt[f"B{R_CUM}"].font=font(9,True,INK); gt[f"C{R_CUM}"]="end of each kit week"; gt[f"C{R_CUM}"].font=font(8,c=GREY)
gt[f"B{R_STAT}"]="Kit coverage"; gt[f"B{R_STAT}"].font=font(9,True,INK); gt[f"C{R_STAT}"]="vs current supplied totals"; gt[f"C{R_STAT}"].font=font(8,c=GREY)
GKITS=(f'IF(COUNTIF({A_C},{GS})=0,"",IF(COUNTIFS({A_C},{GS},{A_S},"")+COUNTIFS({A_C},{GS},tblAssemblies[Check (calc)],"DUPLICATE*")+COUNTIFS({A_C},{GS},tblAssemblies[Check (calc)],"Unknown*")>0,"",MINIFS({A_S},{A_C},{GS})))')
gt["D18"]="="+X(GKITS); gt["D18"].font=font(9,True,OK); gt["D18"].alignment=C
gt["D17"]=""
for k in range(GW):
    col=CL(G0+k)
    gt[f"{col}{R_CUM}"]=f'=SUMIFS({P_Q},{P_C},{GS},{P_D},"<"&({col}${HD}+7*({KLg}+1)))'
    gt[f"{col}{R_CUM}"].font=font(8,c=GREY); gt[f"{col}{R_CUM}"].alignment=C
    gt[f"{col}{R_STAT}"]=(f'=IF({col}{R_K}="","",IF(NOT(ISNUMBER($D${R_STAT})),"setup",IF($D${R_STAT}>={col}{R_CUM},"covered",IF({col}${HD}<=ThisWeek,"NOT covered","not yet"))))')
    gt[f"{col}{R_STAT}"].font=font(7,True); gt[f"{col}{R_STAT}"].alignment=C
gt[f"D{R_CUM}"]="kits ▸"; gt[f"D{R_CUM}"].font=font(7,c=GREY)
gt["B19"]="Complete kits supplied now"; gt["B19"].font=font(8,c=GREY,i=True); gt["C19"]='=IF(ISNUMBER(D18),D18&" complete kits (current total, not dated)","Assembly setup required / incomplete")'; gt["C19"].font=font(8,True,OK)
gt.merge_cells("C19:P19")
# conditional formatting
kr=f"E{R_K}:{GL}{R_K}"
gt.conditional_formatting.add(kr,FormulaRule(formula=[f'AND(E{R_K}<>"",E{R_STAT}="covered")'],fill=fill(OK_S),font=Font(name=F,bold=True,color=OK)))
gt.conditional_formatting.add(kr,FormulaRule(formula=[f'AND(E{R_K}<>"",E{R_STAT}="NOT covered")'],fill=fill(CRIT),font=Font(name=F,bold=True,color=WHITE)))
gt.conditional_formatting.add(kr,FormulaRule(formula=[f'E{R_K}<>""'],fill=fill(ROSE_S),font=Font(name=F,bold=True,color=ROSE)))
gt.conditional_formatting.add(f"E{R_A}:{GL}{R_A}",FormulaRule(formula=[f'E{R_A}<>""'],fill=fill(WARN_S),font=Font(name=F,bold=True,color=WARN)))
gt.conditional_formatting.add(f"E{R_H}:{GL}{R_H}",FormulaRule(formula=[f'E{R_H}<>""'],fill=fill(INFO_S),font=Font(name=F,bold=True,color=INFO)))
sr=f"E{R_STAT}:{GL}{R_STAT}"
gt.conditional_formatting.add(sr,FormulaRule(formula=[f'E{R_STAT}="covered"'],font=Font(name=F,size=7,bold=True,color=OK)))
gt.conditional_formatting.add(sr,FormulaRule(formula=[f'E{R_STAT}="NOT covered"'],font=Font(name=F,size=7,bold=True,color=CRIT)))
# current-week line (left+right red border, very pale fill) across header..status rows; matched on week-start date
red=Side(style="thin",color=CRIT)
gt.conditional_formatting.add(f"E{HW}:{GL}{HW}",FormulaRule(formula=[f'E${HD}=ThisWeek'],fill=fill(CRIT),font=Font(name=F,bold=True,color=WHITE),border=Border(left=red,right=red)))
gt.conditional_formatting.add(f"E{HY}:{GL}{R_STAT+1}",FormulaRule(formula=[f'E${HD}=ThisWeek'],border=Border(left=red,right=red),fill=fill("FFF5F5")))
gt.freeze_panes=gt[f"E{HD+1}"]
gt.protection.sheet=True; gt.protection.formatColumns=False
gt.page_setup.orientation="landscape"; gt.page_setup.fitToWidth=1; gt.page_setup.fitToHeight=1; gt.sheet_properties.pageSetUpPr.fitToPage=True

# =====================================================================
# START HERE
# =====================================================================
st=start; st.sheet_view.showGridLines=False
st.column_dimensions["A"].width=2; st.column_dimensions["B"].width=30; st.column_dimensions["C"].width=110
band(st,range(2,4),(1,2))
st["B1"]="R5B AMP KITTING CONTROL — START HERE"; st["B1"].font=font(20,True,WHITE); st.row_dimensions[1].height=34
st["B2"]="Microsoft 365 workbook. No macros, no external connections. Yellow = you type. Grey = calculated, do not type."; st["B2"].font=font(10,c=MUTE)
sections=[
("EVERY WEEK — UPDATE SUPPLIED TOTALS",[
 ("Where","'Assemblies & Totals'. Find the row for the contract and assembly code."),
 ("Rule","Enter the TOTAL supplied to date, not this week's additions. If the existing total is 26 and another 4 are supplied, enter 30."),
 ("Then fill","Updated on (date) · Updated by (your name) · Source reference (e.g. the SAP export filename you checked)."),
 ("Worked example","Spool 92RRA00548AAA shows 312. This week 24 more were supplied. Type 336 in 'Cumulative supplied to date', today's date in 'Updated on', your name, and e.g. 'COOIS_E2A_2026-10-12.xlsx'."),
 ("Correcting a mistake","Type the correct total (it may be lower) and write why in 'Correction note', e.g. '-3: duplicate WO counted on 05-Oct'. The workbook does not keep old values, so the note is the record."),
]),
("READING THE DASHBOARD",[
 ("Pick a contract","Yellow cell at the top of 'Dashboard'. The list grows automatically when contracts are added."),
 ("Complete repeater kits released","One complete repeater kit means every required assembly has been supplied. It does not mean the repeater has finished production. It is the LOWEST number of complete sets across the contract's assemblies; the limiting row is red in the list."),
 ("Kits due by this week","Repeaters whose kit-release week (final build week minus the kit lead) is this week or earlier. Compared with current plan, not the original commitment."),
 ("Ahead / behind current plan","Complete kits minus kits due by this week."),
 ("Kits due over the next four weeks","Repeaters whose kit-release week falls in the next four weeks."),
 ("Target","Yellow cell. Blank = kits due by this week + next four weeks. Type a number to plan a specific batch. 'Additional needed' = max(0, target × qty per repeater − supplied)."),
 ("Supplied totals line","Shows the latest and oldest update dates, mixed dates, missing dates, and OUT OF DATE when the oldest is older than the interval on 'Contracts' (cell C4)."),
 ("Next action","One concrete step taken from the data (setup → update totals → shortage → covered). It never states production status, availability or completion dates."),
 ("Gantt","Pick a contract. Weekly quantities for kit release, AMPS build and final build (scheduled). Red line = this week. 'Kit coverage' compares cumulative kits due with current complete kits. Plan rows outside the 52-week view are counted above the grid."),
]),
("CONTRACT SETUP (once per contract)",[
 ("1. Contracts","Add a row: unique Contract ID, friendly name, project, top-level ++ code, amps per repeater, total repeaters required, kit-release lead (default 2), AMPS-build lead (default 1), plan revision/date/updater, notes."),
 ("2. Assemblies & Totals","One row per required assembly code from the ZMRP breakdown: contract, group (LASER ASSY…), code, description, qty per repeater, cumulative supplied (0 if none), update details. The same code in two contracts = two rows."),
 ("3. Weekly Plan","One row per final build week: contract, Monday date, repeaters planned."),
 ("Checks","Each input sheet has a Check column. Anything other than OK needs fixing. 'Assembly setup required' = no assemblies yet."),
]),
("PLAN MAINTENANCE",[
 ("Move a batch","Change the Monday date on its row."),
 ("Change quantity","Change 'Repeaters planned'."),
 ("Split","Reduce the quantity on the row and add a new row for the other week."),
 ("Cancel","Delete the row (or set the quantity to 0)."),
 ("Replace a contract's whole plan","Filter 'Contract' on 'Weekly Plan', select and delete ALL of that contract's rows (right-click → Delete → Table Rows), clear the filter, paste the new rows at the bottom of the table. Do not keep old and new rows together. Update Plan revision / updated on / updated by on 'Contracts'. Other contracts are not affected."),
 ("Dates","Always the MONDAY of the final build (housing) week. ISO week labels are calculated, so year boundaries (W52/W53 → W01) work automatically."),
]),
("SOURCES, FILES AND BACKUPS",[
 ("Verified supplied totals","Production controller / kitting lead, from SAP (COOIS for released orders, CO41 for release dates, ZMRP for the assembly breakdown). Record the export filename in 'Source reference'."),
 ("Plan","Production planning (PLAN / ASP export). Record its name in 'Plan revision' on 'Contracts'."),
 ("SAP exports","Keep full exports in the company folder for SAP exports — never paste them into this workbook (it keeps it fast)."),
 ("History archive","'R5B-AMP-Release-Log-Archive.xlsx' holds the old Release Log (182 lines) and the migration reconciliation. It is read-only history; this workbook does not use it."),
 ("Backups","Save a dated copy before replacing a plan or doing large updates (e.g. R5B-AMP-Kitting-Control_2026-10-05.xlsx) in the team backup folder."),
 ("Why no import","A this-week-only SAP export cannot give the cumulative total without a verified opening balance, so totals are typed from a verified cumulative figure."),
]),
("IF A FORMULA IS OVERWRITTEN",[
 ("Grey table columns","Copy a grey cell from the row above and paste it into the broken cell; or type the formula once in any cell of the column and Excel refills the whole column. Formulas are listed below."),
 ("Complete sets (calc)",'Type: =IF(AND(ISNUMBER([@[Qty per repeater]]),[@[Qty per repeater]]>0,[@[Qty per repeater]]=INT([@[Qty per repeater]])),INT(IF(ISNUMBER([@[Cumulative supplied to date]]),[@[Cumulative supplied to date]],0)/[@[Qty per repeater]]),"")'),
 ("Check columns","Copy from the row above (same table)."),
 ("ISO week (calc)",'Type: =IF(ISNUMBER([@[Final build week (Monday)]]),"W"&TEXT(ISOWEEKNUM([@[Final build week (Monday)]]),"00")&"-"&YEAR([@[Final build week (Monday)]]-WEEKDAY([@[Final build week (Monday)]],2)+4),"")'),
 ("Dashboard / Gantt","These sheets are protected (no password) so formulas cannot be overwritten by accident. Review → Unprotect Sheet only for maintenance, then protect again."),
]),
]
r=4
for title,items in sections:
    st[f"B{r}"]=title; st[f"B{r}"].font=font(12,True,WHITE); st[f"B{r}"].fill=fill(ROSE); st[f"C{r}"].fill=fill(ROSE); r+=1
    for k,v in items:
        st[f"B{r}"]=k; st[f"B{r}"].font=font(10,True); st[f"B{r}"].alignment=LT
        st[f"C{r}"]=v; st[f"C{r}"].font=font(10); st[f"C{r}"].alignment=LT
        st.row_dimensions[r].height=max(18,15*(1+len(v)//120)); r+=1
    r+=1
st.protection.sheet=True

# order & finish
wb.calculation=CalcProperties(fullCalcOnLoad=True)
wb.active=1
wb.save(OUT)
json.dump(DA_CELLS,open(OUT+".da.json","w"))

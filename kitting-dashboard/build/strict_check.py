"""Excel-strict structural checks LibreOffice does not enforce."""
import sys,zipfile,re
from openpyxl import load_workbook
from xml.etree import ElementTree as ET
fn=sys.argv[1]; problems=[]
wb=load_workbook(fn)
z=zipfile.ZipFile(fn); ns={"m":"http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
allnames=set()
for ws in wb.worksheets:
    for tname in list(ws.tables.keys()):
        t=ws.tables[tname]
        if tname.lower() in allnames: problems.append(f"duplicate table name {tname}")
        allnames.add(tname.lower())
        m=re.match(r"([A-Z]+)(\d+):([A-Z]+)(\d+)",t.ref); hr=int(m.group(2))
        from openpyxl.utils import column_index_from_string as ci
        c1,c2=ci(m.group(1)),ci(m.group(3))
        hdr=[ws.cell(hr,c).value for c in range(c1,c2+1)]
        names=[tc.name for tc in t.tableColumns]
        if len(names)!=c2-c1+1: problems.append(f"{tname}: column count {len(names)} != ref width {c2-c1+1}")
        for h,n in zip(hdr,names):
            if not isinstance(h,str): problems.append(f"{tname}: header '{h}' is not text")
            elif h!=n: problems.append(f"{tname}: header cell '{h}' != tableColumn name '{n}'")
        if len({n.lower() for n in names})!=len(names): problems.append(f"{tname}: duplicate column names")
        ids=[tc.id for tc in t.tableColumns]
        if len(set(ids))!=len(ids): problems.append(f"{tname}: duplicate column ids")
        # merged cells overlapping table
        for mr in ws.merged_cells.ranges:
            if not (mr.max_row<hr or mr.min_row>int(m.group(4)) or mr.max_col<c1 or mr.min_col>c2): problems.append(f"{tname}: merged cells {mr} overlap table")
        # structured refs used anywhere must name existing columns
# check every structured reference in the workbook resolves to a real table column
tables={}
for ws in wb.worksheets:
    for tname in list(ws.tables.keys()): t=ws.tables[tname]; tables[tname.lower()]={tc.name.lower() for tc in t.tableColumns}
text=""
for n in z.namelist():
    if n.endswith(".xml"): text+=z.read(n).decode("utf8","ignore")
for tb,inner in re.findall(r"(tbl\w+)\[((?:\[[^\]]*\]|[^\]\[])*)\]",text):
    cols=re.findall(r"\[([^\]#][^\]]*)\]",inner) or ([inner] if inner and not inner.startswith("[") else [])
    for c in cols:
        c=c.replace("&amp;","&")
        if tb.lower() not in tables: problems.append(f"unknown table {tb}")
        elif c.lower() not in tables[tb.lower()]: problems.append(f"{tb}: unknown column [{c}]")
print("STRICT CHECK:", "PASS" if not problems else "FAIL"); [print(" -",p) for p in sorted(set(problems))[:40]]

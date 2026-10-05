import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule, DataBarRule, CellIsRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.comments import Comment

OUT = sys.argv[1]

import json
CFG=json.load(open(sys.argv[2]))
SAMPLE=len(sys.argv)>3
REPSEL=(int(sys.argv[4]) if len(sys.argv)>4 else "All")
PLAN=[tuple(x) for x in CFG["plan"]]
SEQS=[]
for p in PLAN:
    if p[0] not in SEQS: SEQS.append(p[0])
NP=len(PLAN); KL=4+NP
NREP=CFG["repeaters"]; RREL=CFG["reps_to_release"]
REP_LIST='"All,'+",".join(str(i) for i in range(1,NREP+1))+'"'
D=CFG["details"]
REPWORD=f"{RREL} repeater"+("s" if RREL>1 else "")
STATUSES = ["Planned","Released","Kitted","Closed"]

NAVY="141A26"; INK="1E2433"; ROSE="B0245E"; ROSE_SOFT="FBE8EF"; GREY="5D6576"; LINE="D9DDE5"; PANEL="F3F4F7"; WHITE="FFFFFF"
OK="1F8A5B"; OK_S="E3F4EC"; WARN="A8660B"; WARN_S="FBF0DC"; CRIT="C0392B"; CRIT_S="FBE6E3"; INFO="3558B8"; INFO_S="E6ECFA"
F="Arial"
def font(sz=11,b=False,c=INK,i=False): return Font(name=F,size=sz,bold=b,color=c,italic=i)
def fill(c): return PatternFill("solid",start_color=c,end_color=c)
thin=Side(style="thin",color=LINE)
BOX=Border(top=thin,bottom=thin,left=thin,right=thin)
BOT=Border(bottom=thin)
C=Alignment(horizontal="center",vertical="center",wrap_text=True)
L=Alignment(horizontal="left",vertical="center",wrap_text=True)
R=Alignment(horizontal="right",vertical="center")
BLUE="0000FF"

wb=Workbook()

# ---------------- Kitting Plan ----------------
kp=wb.active; kp.title="Kitting Plan"
kp.sheet_view.showGridLines=False
kp["A1"]="KITTING PLAN — MASTER DATA"; kp["A1"].font=font(16,True,NAVY)
kp["A2"]=f"Source: kitting plan table supplied by the user ({D['contract']} · {D['program']} · {D['erp']} · {D['config']} · {D['week']}). Blue cells are inputs; edit them when the plan changes."
kp["A2"].font=font(10,c=GREY,i=True)
hdr=["Contract","Sequence","Assy Code","Assy Description","Release","Repeater No.","No. of Rep to Rel","Default Qty","Check"]
for j,h in enumerate(hdr,1):
    c=kp.cell(4,j,h); c.font=font(10,True,WHITE); c.fill=fill(NAVY); c.alignment=C; c.border=BOX
for i,(s,code,d,rel,de) in enumerate(PLAN):
    r=5+i
    vals=[D["contract"],s,code,d,rel,f"1-{NREP}",RREL,de]
    for j,v in enumerate(vals,1):
        c=kp.cell(r,j,v); c.font=font(10,c=BLUE); c.border=BOX
        c.alignment=L if j in (2,3,4) else C
    c=kp.cell(r,9,f'=IF(E{r}=H{r}*G{r},"OK","CHECK: release ≠ default × reps ("&H{r}*G{r}&")")')
    c.font=font(10); c.border=BOX; c.alignment=L
kp.conditional_formatting.add(f"I5:I{KL}",FormulaRule(formula=['LEFT($I5,5)="CHECK"'],fill=fill(WARN_S),font=Font(name=F,color=WARN,bold=True)))
for col,w in zip("ABCDEFGHI",[10,15,18,44,10,13,16,12,36]): kp.column_dimensions[col].width=w
kp.freeze_panes="A5"
kp[f"A{KL+2}"]="Contract details shown on the Dashboard"; kp[f"A{KL+2}"].font=font(11,True,NAVY)
details=[("Contract",D["contract"]),("Program",D["program"]),("ERP item",D["erp"]),("Config",D["config"]),("Kitting week",D["week"]),("Repeaters",f"1-{NREP} · {RREL} to release")]
for i,(k,v) in enumerate(details):
    kp.cell(KL+3+i,1,k).font=font(10,c=GREY)
    c=kp.cell(KL+3+i,2,v); c.font=font(10,c=BLUE)

KP_CODE=f"'Kitting Plan'!$C$5:$C${KL}"; KP_SEQ=f"'Kitting Plan'!$B$5:$B${KL}"; KP_DESC=f"'Kitting Plan'!$D$5:$D${KL}"
KP_REL=f"'Kitting Plan'!$E$5:$E${KL}"; KP_DEF=f"'Kitting Plan'!$H$5:$H${KL}"; KP_REPS=f"'Kitting Plan'!$G$5:$G${KL}"

# ---------------- Work Orders ----------------
wo=wb.create_sheet("Work Orders")
wo.sheet_view.showGridLines=False
N0,N1=6,505
wo["A1"]="WORK ORDERS"; wo["A1"].font=font(16,True,NAVY)
wo["A2"]="Type the WO number and pick the assy code. Quantity fills in automatically from the kitting plan (release qty, or default qty when a single repeater is chosen)."
wo["A2"].font=font(10,c=GREY,i=True)
wo["A3"]="Blue columns = you type. Grey columns = automatic. Use 'Qty Override' only when the WO quantity differs from the plan."
wo["A3"].font=font(10,c=GREY,i=True)
whdr=["WO No.","Assy Code","Description","Sequence","Repeater","Listed Qty","Qty Override","WO Qty","Status","Release Date","Remarks","Flag"]
inputs={1,2,5,7,9,10,11}
for j,h in enumerate(whdr,1):
    c=wo.cell(5,j,h); c.font=font(10,True,WHITE); c.fill=fill(ROSE if j in inputs else NAVY); c.alignment=C; c.border=BOX
for r in range(N0,N1+1):
    wo.cell(r,3,f'=IF(B{r}="","",IFERROR(INDEX({KP_DESC},MATCH(B{r},{KP_CODE},0)),"Unknown code"))')
    wo.cell(r,4,f'=IF(B{r}="","",IFERROR(INDEX({KP_SEQ},MATCH(B{r},{KP_CODE},0)),""))')
    wo.cell(r,6,f'=IF(B{r}="","",IFERROR(IF(OR(E{r}="",E{r}="All"),INDEX({KP_REL},MATCH(B{r},{KP_CODE},0)),INDEX({KP_DEF},MATCH(B{r},{KP_CODE},0))),""))')
    wo.cell(r,8,f'=IF(B{r}="","",IF(G{r}="",F{r},G{r}))')
    wo.cell(r,12,f'=IF(A{r}="","",IF(B{r}="","Missing code",IF(C{r}="Unknown code","Unknown code",IF(COUNTIF($A${N0}:$A${N1},A{r})>1,"Duplicate WO",""))))')
    for j in range(1,13):
        c=wo.cell(r,j); c.border=BOX
        c.font=font(10,c=BLUE if j in inputs else GREY)
        c.alignment=C if j in (5,6,7,8,9,10) else L
    wo.cell(r,8).font=font(10,True,INK)
    wo.cell(r,10).number_format="dd-mmm-yy"
for col,w in zip("ABCDEFGHIJKL",[18,18,42,15,11,11,13,10,12,13,28,16]): wo.column_dimensions[col].width=w
wo.freeze_panes="C6"
wo.auto_filter.ref=f"A5:L{N1}"
dv_code=DataValidation(type="list",formula1=f"=Lists!$A$1:$A${NP}",allow_blank=True,showErrorMessage=True,errorTitle="Unknown code",error="Pick an assy code from the kitting plan.")
dv_rep=DataValidation(type="list",formula1=REP_LIST,allow_blank=True)
dv_st=DataValidation(type="list",formula1='"Planned,Released,Kitted,Closed"',allow_blank=True)
dv_q=DataValidation(type="whole",operator="greaterThan",formula1="0",allow_blank=True,showErrorMessage=True,error="Enter a whole number above 0, or leave blank to use the listed qty.")
for dv,col in ((dv_code,"B"),(dv_rep,"E"),(dv_st,"I"),(dv_q,"G")):
    wo.add_data_validation(dv); dv.add(f"{col}{N0}:{col}{N1}")
rng=f"A{N0}:L{N1}"
wo.conditional_formatting.add(f"L{N0}:L{N1}",FormulaRule(formula=[f'$L{N0}<>""'],fill=fill(CRIT_S),font=Font(name=F,color=CRIT,bold=True)))
for st,fc,bc in (("Released",INFO,INFO_S),("Kitted",WARN,WARN_S),("Closed",OK,OK_S)):
    wo.conditional_formatting.add(f"I{N0}:I{N1}",FormulaRule(formula=[f'$I{N0}="{st}"'],fill=fill(bc),font=Font(name=F,color=fc,bold=True)))
wo["G4"]="Leave blank to use Listed Qty"; wo["G4"].font=font(8,c=GREY,i=True); wo["G4"].alignment=C
wo["E4"]="Blank = All"; wo["E4"].font=font(8,c=GREY,i=True); wo["E4"].alignment=C
wo["I4"]="Blank = Planned"; wo["I4"].font=font(8,c=GREY,i=True); wo["I4"].alignment=C

# Lists (hidden)
ls=wb.create_sheet("Lists")
for i in range(NP): ls.cell(i+1,1,f"='Kitting Plan'!C{5+i}")
ls.sheet_state="hidden"

W_A=f"'Work Orders'!$A${N0}:$A${N1}"; W_B=f"'Work Orders'!$B${N0}:$B${N1}"; W_E=f"'Work Orders'!$E${N0}:$E${N1}"
W_H=f"'Work Orders'!$H${N0}:$H${N1}"; W_I=f"'Work Orders'!$I${N0}:$I${N1}"; W_L=f"'Work Orders'!$L${N0}:$L${N1}"

# ---------------- Dashboard ----------------
ds=wb.create_sheet("Dashboard",0)
ds.sheet_view.showGridLines=False
ds.sheet_view.zoomScale=80
widths={"A":2,"B":15,"C":18,"D":44,"E":11,"F":11,"G":12,"H":11,"I":15,"J":13,"K":8,"L":12,"M":34,"N":24,"O":3,"P":3}
for k,v in widths.items(): ds.column_dimensions[k].width=v
for col in "BCDEFGHIJKLMN":
    for r in range(1,5): ds[f"{col}{r}"].fill=fill(NAVY)
ds.row_dimensions[2].height=40; ds.row_dimensions[3].height=22
ds["B2"]=CFG["title"]; ds["B2"].font=font(26,True,WHITE); ds["B2"].alignment=Alignment(vertical="center")
ds["B3"]=f"=\"Contract \"&'Kitting Plan'!B{KL+3}&\"   ·   \"&'Kitting Plan'!B{KL+4}&\"   ·   \"&'Kitting Plan'!B{KL+5}&\"   ·   \"&'Kitting Plan'!B{KL+6}&\"   ·   Repeaters \"&'Kitting Plan'!B{KL+8}"
ds["B3"].font=font(12,c="AEB6C6")
ds["L2"]="VIEW REPEATER ▸"; ds["L2"].font=font(11,True,"AEB6C6"); ds["L2"].alignment=R
ds["M2"]=REPSEL; ds["M2"].font=font(20,True,WHITE); ds["M2"].fill=fill(ROSE); ds["M2"].alignment=C
ds["M3"]=f"Pick All or 1–{NREP}"; ds["M3"].font=font(9,c="AEB6C6",i=True); ds["M3"].alignment=C
ds["N2"]="=IF(M2=\"All\",\"Required = release qty\",\"Required = default qty for rep \"&M2)"; ds["N2"].font=font(10,c="AEB6C6"); ds["N2"].alignment=L
dv_sel=DataValidation(type="list",formula1=REP_LIST,allow_blank=False); ds.add_data_validation(dv_sel); dv_sel.add("M2")
SEL="$M$2"

# Ledger first (rows 21-36) so KPIs can reference it
LH=20; L0=21; L1=20+NP
ds.row_dimensions[19].height=26
ds["B19"]="ASSEMBLY LEDGER"; ds["B19"].font=font(14,True,NAVY)
ds["D19"]="Coverage = WO qty ÷ required qty for the selected view. Repeater views count each 'All' WO split evenly across the repeaters being released."
ds["D19"].font=font(9,c=GREY,i=True)
lhdr=["Sequence","Assy Code","Description","Release","Default / Rep","Required","WO Qty","Coverage","Status","WOs","Closed Qty","Work Order Nos.","Plan Check"]
for j,h in enumerate(lhdr):
    c=ds.cell(LH,2+j,h); c.font=font(10,True,WHITE); c.fill=fill(INK); c.alignment=C; c.border=BOX
ds.row_dimensions[LH].height=30
for i in range(NP):
    r=L0+i; k=5+i
    ds[f"B{r}"]=f"='Kitting Plan'!B{k}"; ds[f"C{r}"]=f"='Kitting Plan'!C{k}"; ds[f"D{r}"]=f"='Kitting Plan'!D{k}"
    ds[f"E{r}"]=f"='Kitting Plan'!E{k}"; ds[f"F{r}"]=f"='Kitting Plan'!H{k}"
    ds[f"G{r}"]=f'=IF({SEL}="All",E{r},F{r})'
    ds[f"H{r}"]=(f'=IF({SEL}="All",SUMIFS({W_H},{W_B},C{r}),'
                 f'SUMIFS({W_H},{W_B},C{r},{W_E},{SEL})+(SUMIFS({W_H},{W_B},C{r},{W_E},"All")+SUMIFS({W_H},{W_B},C{r},{W_E},""))/\'Kitting Plan\'!G{k})')
    ds[f"I{r}"]=f'=IF(G{r}=0,0,H{r}/G{r})'
    ds[f"J{r}"]=f'=IF(H{r}=0,"NO WO",IF(H{r}>G{r}+0.0001,"OVER",IF(H{r}>=G{r}-0.0001,"COVERED","SHORT "&ROUND(G{r}-H{r},1))))'
    ds[f"K{r}"]=f'=COUNTIFS({W_B},C{r},{W_A},"?*")'
    ds[f"L{r}"]=f'=SUMIFS({W_H},{W_B},C{r},{W_I},"Closed")'
    ds[f"M{r}"]=ArrayFormula(f"M{r}",f'=_xlfn.TEXTJOIN(", ",TRUE,IF(({W_B}=C{r})*({W_A}<>""),{W_A},""))')
    ds[f"N{r}"]=f"='Kitting Plan'!I{k}"
    ds[f"P{r}"]=f"=MIN(H{r},G{r})"  # helper: capped coverage
    for col in "BCDEFGHIJKLMN":
        c=ds[f"{col}{r}"]; c.border=Border(bottom=thin); c.font=font(11)
        c.alignment=C if col in "EFGHIJKL" else L
    ds[f"C{r}"].font=font(11,True); ds[f"G{r}"].font=font(12,True); ds[f"H{r}"].font=font(12,True)
    ds[f"D{r}"].font=font(10,c=GREY); ds[f"M{r}"].font=font(10); ds[f"N{r}"].font=font(9,c=GREY)
    ds[f"I{r}"].number_format="0%"; ds[f"H{r}"].number_format="#,##0.#;-#,##0.#;0"; ds[f"G{r}"].number_format="#,##0"
    ds.row_dimensions[r].height=24
ds.column_dimensions["P"].hidden=True
LR=f"{L0}:{L1}"
ds.conditional_formatting.add(f"I{L0}:I{L1}",DataBarRule(start_type="num",start_value=0,end_type="num",end_value=1,color=ROSE,showValue=True))
for cond,fc,bc in (('$J{r}="NO WO"',GREY,"ECEEF3"),('LEFT($J{r},5)="SHORT"',WARN,WARN_S),('$J{r}="COVERED"',OK,OK_S),('$J{r}="OVER"',CRIT,CRIT_S)):
    ds.conditional_formatting.add(f"J{L0}:J{L1}",FormulaRule(formula=[cond.format(r=L0)],fill=fill(bc),font=Font(name=F,color=fc,bold=True)))
ds.conditional_formatting.add(f"N{L0}:N{L1}",FormulaRule(formula=[f'LEFT($N{L0},5)="CHECK"'],fill=fill(WARN_S),font=Font(name=F,color=WARN,bold=True,size=9)))
ds.conditional_formatting.add(f"B{L0}:B{L1}",FormulaRule(formula=[f'$B{L0}=$B{L0-1}'],font=Font(name=F,color="C3C8D2")))

# KPI tiles rows 6-8
tiles=[("B","C","ASSEMBLIES",f"=COUNTA(C{L0}:C{L1})",f'="{len(SEQS)} sequences · "&COUNTIF(J{L0}:J{L1},"COVERED")&" fully covered"',"0"),
       ("D","D","REQUIRED QTY",f"=SUM(G{L0}:G{L1})",f'=IF({SEL}="All","units · release across {REPWORD}","units · default qty, repeater "&{SEL})',"#,##0"),
       ("E","G","WO COVERAGE",f"=IF(SUM(G{L0}:G{L1})=0,0,SUM(P{L0}:P{L1})/SUM(G{L0}:G{L1}))",f'=TEXT(SUM(P{L0}:P{L1}),"#,##0")&" of "&TEXT(SUM(G{L0}:G{L1}),"#,##0")&" units on WOs"',"0%"),
       ("H","I","WORK ORDERS",f'=COUNTIF({W_A},"?*")',f'=COUNTIFS({W_A},"?*",{W_I},"Closed")&" closed · "&COUNTIF(J{L0}:J{L1},"NO WO")&" codes with no WO"',"0"),
       ("J","K","SHORT / OVER",f'=COUNTIF(J{L0}:J{L1},"SHORT*")+COUNTIF(J{L0}:J{L1},"OVER")',f'=COUNTIF(J{L0}:J{L1},"SHORT*")&" short · "&COUNTIF(J{L0}:J{L1},"OVER")&" over"',"0"),
       ("L","N","DATA FLAGS",f'=COUNTIF(N{L0}:N{L1},"CHECK*")+COUNTIF({W_L},"?*")',f'=COUNTIF(N{L0}:N{L1},"CHECK*")&" plan qty checks · "&COUNTIF({W_L},"?*")&" WO entry flags"',"0")]
ds.row_dimensions[6].height=22; ds.row_dimensions[7].height=46; ds.row_dimensions[8].height=22
for a,b,lab,val,sub,nf in tiles:
    for r in (6,7,8):
        if a!=b: ds.merge_cells(f"{a}{r}:{b}{r}")
        for col in range(ord(a),ord(b)+1):
            c=ds[f"{chr(col)}{r}"]; c.fill=fill(PANEL)
            c.border=Border(left=Side(style="thick",color=ROSE) if chr(col)==a else None)
    ds[f"{a}6"]=lab; ds[f"{a}6"].font=font(10,True,GREY); ds[f"{a}6"].alignment=Alignment(horizontal="left",vertical="bottom",indent=1)
    ds[f"{a}7"]=val; ds[f"{a}7"].font=font(30,True,NAVY); ds[f"{a}7"].alignment=Alignment(horizontal="left",vertical="center",indent=1); ds[f"{a}7"].number_format=nf
    ds[f"{a}8"]=sub; ds[f"{a}8"].font=font(9,c=GREY); ds[f"{a}8"].alignment=Alignment(horizontal="left",vertical="top",indent=1,wrap_text=True)
ds.conditional_formatting.add("L7",CellIsRule(operator="greaterThan",formula=["0"],font=Font(name=F,color=WARN,bold=True,size=30)))
ds.conditional_formatting.add("J7",CellIsRule(operator="greaterThan",formula=["0"],font=Font(name=F,color=WARN,bold=True,size=30)))

# Sequence table rows 10-17
ds["B10"]="COVERAGE BY SEQUENCE"; ds["B10"].font=font(14,True,NAVY)
for j,h in enumerate(["Sequence","Codes","Coverage","Required","WO Qty","%"]):
    c=ds.cell(11,2+j,h); c.font=font(10,True,WHITE); c.fill=fill(INK); c.alignment=C
for i,s in enumerate(SEQS):
    r=12+i
    ds[f"B{r}"]=s; ds[f"C{r}"]=f'=COUNTIF($B${L0}:$B${L1},B{r})'
    ds[f"E{r}"]=f'=SUMIFS($G${L0}:$G${L1},$B${L0}:$B${L1},B{r})'
    ds[f"F{r}"]=f'=SUMIFS($P${L0}:$P${L1},$B${L0}:$B${L1},B{r})'
    ds[f"G{r}"]=f'=IF(E{r}=0,0,F{r}/E{r})'
    ds[f"D{r}"]=f"=G{r}"
    for col in "BCDEFG":
        c=ds[f"{col}{r}"]; c.border=BOT; c.font=font(11,col in "BG"); c.alignment=L if col=="B" else C
    ds[f"G{r}"].number_format="0%"; ds[f"F{r}"].number_format="#,##0.#;-#,##0.#;0"; ds[f"D{r}"].number_format=";;;"
    ds.row_dimensions[r].height=22
ds.conditional_formatting.add("D12:D17",DataBarRule(start_type="num",start_value=0,end_type="num",end_value=1,color=ROSE,showValue=False))

# Pipeline I..N rows 10-15
ds["I10"]="WORK-ORDER PIPELINE"; ds["I10"].font=font(14,True,NAVY)
for j,h in enumerate(["Status","WOs","Qty","Share of WOs"]):
    c=ds.cell(11,9+j,h); c.font=font(10,True,WHITE); c.fill=fill(INK); c.alignment=C
ds.merge_cells("L11:N11")
stc={"Planned":(GREY,"ECEEF3"),"Released":(INFO,INFO_S),"Kitted":(WARN,WARN_S),"Closed":(OK,OK_S)}
for i,st in enumerate(STATUSES):
    r=12+i
    ds[f"I{r}"]=st
    if st=="Planned":
        ds[f"J{r}"]=f'=COUNTIFS({W_A},"?*",{W_I},"Planned")+COUNTIFS({W_A},"?*",{W_I},"")'
        ds[f"K{r}"]=f'=SUMIFS({W_H},{W_A},"?*",{W_I},"Planned")+SUMIFS({W_H},{W_A},"?*",{W_I},"")'
    else:
        ds[f"J{r}"]=f'=COUNTIFS({W_A},"?*",{W_I},"{st}")'
        ds[f"K{r}"]=f'=SUMIFS({W_H},{W_A},"?*",{W_I},"{st}")'
    ds[f"L{r}"]=f'=IF($H$7=0,0,J{r}/$H$7)'
    ds.merge_cells(f"L{r}:N{r}")
    fc,bc=stc[st]
    ds[f"I{r}"].font=font(11,True,fc); ds[f"I{r}"].fill=fill(bc)
    for col in "IJKLMN":
        c=ds[f"{col}{r}"]; c.border=BOT; c.alignment=C if col!="I" else L
    ds[f"J{r}"].font=font(14,True,NAVY); ds[f"K{r}"].font=font(11,c=GREY); ds[f"L{r}"].number_format="0%"; ds[f"L{r}"].font=font(10,c=GREY)
ds.conditional_formatting.add("L12:L15",DataBarRule(start_type="num",start_value=0,end_type="num",end_value=1,color="8A91A0",showValue=True))
ds["I16"]="Total"; ds["I16"].font=font(11,True)
ds["J16"]="=SUM(J12:J15)"; ds["K16"]="=SUM(K12:K15)"
for col in "JK": ds[f"{col}16"].font=font(11,True); ds[f"{col}16"].alignment=C
ds.merge_cells("I17:N17")
ds["I17"]=(f'=IF(COUNTIF(N{L0}:N{L1},"CHECK*")=0,"✓ Every release qty = default qty × repeaters",'
           f'"⚠ "&COUNTIF(N{L0}:N{L1},"CHECK*")&" code(s): release qty ≠ default × repeaters — see Plan Check column")')
ds["I17"].font=font(10,True,WARN); ds["I17"].fill=fill(WARN_S); ds["I17"].alignment=L
ds.row_dimensions[17].height=24
ds.freeze_panes="A5"

# Print / projection setup
for sh in (ds,):
    sh.page_setup.orientation="landscape"; sh.page_setup.fitToWidth=1; sh.page_setup.fitToHeight=1
    sh.sheet_properties.pageSetUpPr.fitToPage=True
    sh.print_area=f"A1:N{L1}"
wo.page_setup.orientation="landscape"

# ---------------- How To Use ----------------
hu=wb.create_sheet("How To Use")
hu.sheet_view.showGridLines=False
hu.column_dimensions["A"].width=3; hu.column_dimensions["B"].width=26; hu.column_dimensions["C"].width=100
hu["B2"]="HOW TO USE THIS WORKBOOK"; hu["B2"].font=font(18,True,NAVY)
rows=[(k,v,h) for k,v,h in CFG["howto"]]
r=4
for k,v,head in rows:
    if head:
        r+=1; hu[f"B{r}"]=k; hu[f"B{r}"].font=font(12,True,ROSE); r+=1; continue
    hu[f"B{r}"]=k; hu[f"B{r}"].font=font(11,True); hu[f"B{r}"].alignment=Alignment(vertical="top")
    hu[f"C{r}"]=v; hu[f"C{r}"].font=font(11); hu[f"C{r}"].alignment=Alignment(wrap_text=True,vertical="top")
    hu.row_dimensions[r].height=34
    r+=1

if SAMPLE:
    import random
    data=[tuple(x) for x in CFG["sample"]]
    for i,(a,b,e,g,st) in enumerate(data):
        rr=N0+i; wo[f"A{rr}"]=a; wo[f"B{rr}"]=b; wo[f"E{rr}"]=(int(e) if e.isdigit() else e) or None
        if g: wo[f"G{rr}"]=int(g)
        if st: wo[f"I{rr}"]=st
from openpyxl.workbook.properties import CalcProperties
wb.calculation=CalcProperties(fullCalcOnLoad=True)
wb.active=0
wb.save(OUT)

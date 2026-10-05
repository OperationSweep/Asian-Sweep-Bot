import sys, datetime, subprocess, re, copy
from openpyxl import load_workbook
from openpyxl.worksheet.formula import ArrayFormula
S="/tmp/claude-0/-home-user-Asian-Sweep-Bot/0a9d6962-6e96-5f3a-ae3c-2b9037388c4d/scratchpad"
def load(): return load_workbook(f"{S}/v2_raw.xlsx")
def add_rows(wb, sheet, tname, rows):
    """simulate typing new rows directly under a table: Excel extends the table and fills calculated columns."""
    ws=wb[sheet]; t=ws.tables[tname]
    m=re.match(r"([A-Z]+)(\d+):([A-Z]+)(\d+)",t.ref); c1,r1,c2,r2=m.group(1),int(m.group(2)),m.group(3),int(m.group(4))
    hdr=[c.value for c in ws[r1]][:len(t.tableColumns)]
    for i,row in enumerate(rows):
        r=r2+1+i
        for j,h in enumerate(hdr,1):
            tc=t.tableColumns[j-1]
            if tc.calculatedColumnFormula is not None: ws.cell(r,j).value="="+tc.calculatedColumnFormula.attr_text
            elif h in row: ws.cell(r,j).value=row[h]
    t.ref=f"{c1}{r1}:{c2}{r2+len(rows)}"
    if t.autoFilter is not None: t.autoFilter.ref=t.ref
def find_rows(wb, sheet, tname, pred):
    ws=wb[sheet]; t=ws.tables[tname]; m=re.match(r"[A-Z]+(\d+):[A-Z]+(\d+)",t.ref); r1,r2=int(m.group(1)),int(m.group(2))
    hdr=[c.value for c in ws[r1]][:len(t.tableColumns)]
    out=[]
    for r in range(r1+1,r2+1):
        rec={h:ws.cell(r,j+1).value for j,h in enumerate(hdr)}
        if pred(rec): out.append(r)
    return out,hdr
def setcell(wb,sheet,tname,r,col,val):
    ws=wb[sheet]; _,hdr=find_rows(wb,sheet,tname,lambda x:False); ws.cell(r,hdr.index(col)+1).value=val
def delete_row(wb,sheet,tname,r):
    ws=wb[sheet]; t=ws.tables[tname]; m=re.match(r"([A-Z]+)(\d+):([A-Z]+)(\d+)",t.ref); c1,r1,c2,r2=m.group(1),int(m.group(2)),m.group(3),int(m.group(4))
    ncol=len(t.tableColumns)
    for rr in range(r,r2):
        for j in range(1,ncol+1): ws.cell(rr,j).value=ws.cell(rr+1,j).value
    for j in range(1,ncol+1): ws.cell(r2,j).value=None
    t.ref=f"{c1}{r1}:{c2}{r2-1}"
    if t.autoFilter is not None: t.autoFilter.ref=t.ref
def run(wb,name):
    p=f"{S}/{name}.xlsx"; wb.save(p)
    out=subprocess.run([f"{S}/chk.sh",p],capture_output=True,text=True).stdout.strip()
    return load_workbook(out,data_only=True)
def dash(v):
    d=v["Dashboard"]; P={d[f"O{r}"].value:d[f"P{r}"].value for r in range(1,19)}
    lst=[]
    for r in range(21,90):
        row=[d.cell(r,c).value for c in range(2,13)]
        if row[0] in (None,"#N/A") : break
        lst.append(row)
    return d,P,lst
def errs(v):
    e=[]
    for ws in v.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value,str) and re.match(r"#(N/A|VALUE!|REF!|NAME\?|DIV/0!|NUM!|NULL!)",c.value) and not (ws.title=="Dashboard" and c.row>=21 and c.value=="#N/A"):
                    e.append(f"{ws.title}!{c.coordinate}={c.value}")
    return e

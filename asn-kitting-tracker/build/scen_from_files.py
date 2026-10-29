"""Turn the real ASP / COOIS files into a scenario JSON (dates as Excel serials) so Excel and the JS engine get identical inputs."""
import sys, json, datetime, os
from openpyxl import load_workbook
src = sys.argv[1]; out = sys.argv[2]
def ser(v):
    if isinstance(v, datetime.datetime): return (v - datetime.datetime(1899, 12, 30)).days
    if isinstance(v, datetime.date): return (v - datetime.date(1899, 12, 30)).days
    if isinstance(v, datetime.time): return None
    if isinstance(v, str) and v.startswith("="): return None
    return v
def rows(path, sheet=None, ncols=None):
    ws = load_workbook(path)[sheet] if sheet else load_workbook(path).worksheets[0]
    r = [[ser(c.value) for c in row[:ncols] if True] for row in ws.iter_rows()]
    while r and all(v is None for v in r[-1]): r.pop()
    return r
a = rows(os.path.join(src, "asp.xlsx")); c = rows(os.path.join(src, "coois.xlsx"), "Sheet1", 15)
ref = json.load(open(os.path.join(os.path.dirname(os.path.abspath(__file__)), "reference-data.json")))
json.dump({"asp": {"header": a[0], "rows": a[1:]}, "coois": {"header": c[0], "rows": c[1:]}, "requirements": ref["requirements"],
           "stages": ref["stages"], "settings": {"kitLead": 3}, "log": []}, open(out, "w"))
print(len(a) - 1, "ASP rows", len(c) - 1, "COOIS rows")

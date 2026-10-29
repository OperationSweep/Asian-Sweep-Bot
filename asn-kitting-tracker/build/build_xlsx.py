"""Build the ASN Kitting Progress Tracker workbook (Microsoft 365, no macros).

usage: python3 build_xlsx.py out.xlsx [scenario.json] [--test]
  scenario.json (optional) supplies ASP / COOIS / requirements / log / settings; default = real source files + empty log.
  --test widens dynamic-array ranges so LibreOffice can recalculate them (test copies only, never shipped).
"""
import sys, json, re, os, datetime, zipfile
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.worksheet.table import Table, TableStyleInfo, TableFormula
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.utils import get_column_letter as CL

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = sys.argv[1]
SCEN = next((a for a in sys.argv[2:] if a.endswith(".json")), None)
TEST = "--test" in sys.argv

# ---------------------------------------------------------------- formula helpers
FN = ["FILTER", "SORT", "LET", "HSTACK", "VSTACK", "TAKE", "DROP", "CHOOSECOLS", "SORTBY", "XLOOKUP", "XMATCH", "UNIQUE",
      "SEQUENCE", "MINIFS", "MAXIFS", "ISOWEEKNUM", "TEXTJOIN", "EXPAND", "TOCOL", "IFS"]
PFX = {f: "_xlfn." + f for f in FN}
PFX["FILTER"] = "_xlfn._xlws.FILTER"; PFX["SORT"] = "_xlfn._xlws.SORT"
LETVARS = {'nm', 'nl'}
def X(f):
    f = re.sub(r'(?<![\w.])(' + "|".join(sorted(FN, key=len, reverse=True)) + r')\(', lambda m: PFX[m.group(1)] + "(", f)
    # LET variable names: declared as  LET(name,  or  ,name,  pairs -> prefix _xlpm. on every use
    for v in sorted(LETVARS, key=len, reverse=True):
        f = re.sub(r'(?<![\w.\[\]"])' + v + r'(?![\w(\[])', "_xlpm." + v, f)
    return f

def let(*pairs_and_result):
    """let(('name', expr), ..., result) -> LET formula text; registers names for prefixing."""
    *pairs, res = pairs_and_result
    for n, _ in pairs:
        assert not re.fullmatch(r"[A-Za-z]{1,3}\d+", n), "LET name looks like a cell address: " + n
        LETVARS.add(n)
    return "LET(" + ",".join(f"{n},{e}" for n, e in pairs) + "," + res + ")"

DA = []   # (sheet, cell, test_ref)
def da(ws, cell, formula, test_ref=None):
    f = X(formula); assert not string_leaks(f), (cell, string_leaks(f))
    ws[cell] = ArrayFormula(cell, "=" + f); DA.append((ws.title, cell, test_ref))
def fx(ws, cell, formula):
    f = X(formula); assert not string_leaks(f), (cell, string_leaks(f))
    ws[cell] = "=" + f

def string_leaks(f):
    return [m for m in re.findall(r'"[^"]*"', f) if "_xlpm." in m or "_xlfn." in m]

# ---------------------------------------------------------------- style
NAVY = "17324D"; INK = "1E2433"; TEAL = "0E7C7B"; GREY = "5D6576"; LINE = "D5DAE1"; PANEL = "F2F4F7"; WHITE = "FFFFFF"
OK = "1E7B4F"; OK_S = "D9F0E3"; OK_L = "EEF8F2"; WARN = "9A5B00"; WARN_S = "FCEFD6"; CRIT = "B42318"; CRIT_S = "FBE3E1"
INFO = "2F5AAE"; INFO_S = "E2EAFA"; INPUT = "FFF4BF"; CALC = "E7E9ED"; PASTE = "DDF3E6"
FONT = "Arial"
def font(sz=10, b=False, c=INK, i=False): return Font(name=FONT, size=sz, bold=b, color=c, italic=i)
def fill(c): return PatternFill("solid", start_color=c, end_color=c)
thin = Side(style="thin", color=LINE); box = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top"); CEN = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
UNLOCK = Protection(locked=False)

def title(ws, text, sub=None, width=16):
    ws["B2"] = text; ws["B2"].font = font(18, True, NAVY)
    if sub: ws["B3"] = sub; ws["B3"].font = font(10, False, GREY, True)
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 2

def hdr(ws, row, col, labels, color=NAVY):
    for j, t in enumerate(labels):
        c = ws.cell(row, col + j, t); c.font = font(9, True, WHITE); c.fill = fill(color); c.alignment = CEN; c.border = box

def label(ws, cell, text, sz=10, b=False, c=INK, i=False):
    ws[cell] = text; ws[cell].font = font(sz, b, c, i); ws[cell].alignment = LEFT

def section(ws, cell, text):
    ws[cell] = text; ws[cell].font = font(12, True, TEAL)

# ---------------------------------------------------------------- data
ref = json.load(open(os.path.join(HERE, "reference-data.json")))
if SCEN:
    sc = json.load(open(SCEN))
else:
    sc = {}
REQ = sc.get("requirements", ref["requirements"])
STAGES = sc.get("stages", ref["stages"])
LOGROWS = sc.get("log", [])
SETTINGS = dict(ref["settings"]); SETTINGS.update(sc.get("settings", {}))
SRC = os.environ.get("KIT_SRC", os.path.join(os.path.dirname(HERE), "source"))
BLANK = "--blank" in sys.argv

def read_sheet(path, sheet=None):
    wb = load_workbook(path, data_only=False)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    rows = [[c.value for c in r] for r in ws.iter_rows()]
    fmts = [[c.number_format for c in r] for r in ws.iter_rows(max_row=3)]
    return rows, fmts

if "asp" in sc:
    ASP = [sc["asp"]["header"]] + sc["asp"]["rows"]; ASP_FMT = None
else:
    ASP, ASP_FMT = read_sheet(os.path.join(SRC, "asp.xlsx"))
    if BLANK: ASP = ASP[:1]
if "coois" in sc:
    COOIS = [sc["coois"]["header"]] + sc["coois"]["rows"]; COOIS_FMT = None
else:
    COOIS, COOIS_FMT = read_sheet(os.path.join(SRC, "coois.xlsx"), "Sheet1")
    if BLANK: COOIS = COOIS[:1]
    while COOIS and all(v is None for v in COOIS[-1]): COOIS.pop()

def xl_date(v):
    """scenario JSON stores dates as Excel serial numbers; real files give datetimes"""
    return v

LOG_COLS = ["Entry ID", "Entry date", "Record type", "Variant ID", "Stage", "Assembly code", "REPs kitted", "REP from", "REP to",
            "Order", "Order qty", "Void", "Void reason", "Note"]
LOG_CALC = ["Order key", "Active", "Check"]
REQ_COLS = ["Project", "Variant ID", "Top code", "Fibre pairs", "Stage", "Assembly code", "Assembly description", "Qty per REP",
            "Source", "Verified by / date"]
STAGE_COLS = ["Stage", "Order", "Description"]

ASP_ROWS = 1000; ASP_LASTCOL = "ALL"          # ASP rows 2..1000, week columns Q..ALL (984 weeks)
NPLAN = ASP_ROWS - 1
COOIS_ROWS = 20001                             # COOIS rows 2..20001
RV_LAST = 6 + COOIS_ROWS - 2                   # COOIS Review rows 6..20005
LINE_CAP = 2000

wb = Workbook(); wb.remove(wb.active)
readme = wb.create_sheet("Read Me"); dash = wb.create_sheet("Dashboard"); rv = wb.create_sheet("Repeater View")
rec = wb.create_sheet("Record Kitting"); lg = wb.create_sheet("Kitting Log"); st = wb.create_sheet("Settings")
cr = wb.create_sheet("COOIS Review"); asp = wb.create_sheet("ASP"); co = wb.create_sheet("COOIS"); calc = wb.create_sheet("Calc")
for ws, col in [(dash, NAVY), (rv, NAVY), (rec, TEAL), (lg, TEAL), (st, TEAL), (cr, INFO), (asp, "8A94A6"), (co, "8A94A6"), (calc, "B8BFCC"), (readme, NAVY)]:
    ws.sheet_properties.tabColor = col

def name(n, ref_):
    wb.defined_names[n] = DefinedName(n, attr_text=ref_)

# ---------------------------------------------------------------- Settings
title(st, "Settings and assembly requirements", "Yellow cells and the two tables can be edited. Add rows at the bottom of a table; it grows automatically.")
params = [("Report date (leave blank = today)", None, "dd-mmm-yyyy"), ("Kit lead (weeks before ASP week)", SETTINGS.get("kitLead", 3), "0"),
          ("ASP file loaded", SETTINGS.get("aspFile", "ASP GH SME AUG 2026.xlsx"), "@"), ("ASP loaded on", SETTINGS.get("aspLoaded"), "dd-mmm-yyyy"),
          ("COOIS file loaded", SETTINGS.get("cooisFile", "export.xlsx"), "@"), ("COOIS loaded on", SETTINGS.get("cooisLoaded"), "dd-mmm-yyyy")]
section(st, "B5", "General")
for i, (lab, val, fmt) in enumerate(params):
    r = 6 + i
    label(st, f"B{r}", lab); c = st[f"C{r}"]; c.value = val; c.fill = fill(INPUT); c.border = box; c.number_format = fmt; c.protection = UNLOCK
if SETTINGS.get("reportDate"): st["C6"] = SETTINGS["reportDate"]
name("RptDate", 'IF(Settings!$C$6="",TODAY(),Settings!$C$6)')
name("ThisMon", "RptDate-WEEKDAY(RptDate,2)+1")
name("KitLead", "Settings!$C$7")
st.column_dimensions["B"].width = 34; st.column_dimensions["C"].width = 24

def sref(f, tname):
    """UI shorthand -> file syntax: [@[X]] / [@X] -> tbl[[#This Row],[X]], bare [X] -> tbl[X]"""
    keep = []
    def hold(m):
        keep.append(f"{tname}[[#This Row],[{m.group(1)}]]"); return f"\x00{len(keep) - 1}\x00"
    f = re.sub(r"\[@\[([^\]]+)\]\]", hold, f)
    f = re.sub(r"\[@([^\[\]]+)\]", hold, f)
    f = re.sub(r"(?<![\w\]\[])\[([^\[\]@#][^\[\]]*)\]", lambda m: f"{tname}[{m.group(1)}]", f)
    return re.sub(r"\x00(\d+)\x00", lambda m: keep[int(m.group(1))], f)

def make_table(ws, tname, top_row, left_col, cols, rows, calc=None, style="TableStyleMedium2", widths=None):
    calc = {k: sref(v, tname) for k, v in (calc or {}).items()}
    allc = cols + list(calc)
    hdr(ws, top_row, left_col, allc, NAVY)
    rows = rows or [[None] * len(cols)]
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            c = ws.cell(top_row + 1 + i, left_col + j, v); c.protection = UNLOCK; c.font = font(10)
    colpos = {h: CL(left_col + j) for j, h in enumerate(allc)}
    def rowform(k, r):
        f = X(calc[k])
        if "LET(" in f:   # LibreOffice mis-reads [#This Row] inside LET; same-row cell references work in both programs
            f = re.sub(re.escape(tname) + r"\[\[#This Row\],\[([^\]]+)\]\]", lambda m: f"${colpos[m.group(1)]}{r}", f)
        return f
    for i in range(len(rows)):
        for j, k in enumerate(calc):
            c = ws.cell(top_row + 1 + i, left_col + len(cols) + j, "=" + rowform(k, top_row + 1 + i)); c.fill = fill(CALC); c.font = font(9, False, GREY)
    refs = f"{CL(left_col)}{top_row}:{CL(left_col + len(allc) - 1)}{top_row + len(rows)}"
    t = Table(displayName=tname, ref=refs); t.tableStyleInfo = TableStyleInfo(name=style, showRowStripes=True)
    t._initialise_columns()
    for j, h in enumerate(allc):
        t.tableColumns[j].name = h
        if h in calc: t.tableColumns[j].calculatedColumnFormula = TableFormula(attr_text=rowform(h, top_row + 1))
    ws.add_table(t)
    if widths:
        for j, w in enumerate(widths): ws.column_dimensions[CL(left_col + j)].width = w
    return t

section(st, "B14", "Assembly requirements (one row per assembly line, quantity per ONE repeater)")
label(st, "B15", "Contract (Project), exact variant (ASP Friendly_name), stage, assembly code and quantity per repeater. "
      "A variant with no rows here shows \"Assembly setup required\". Quantities reconstructed from pasted tables are marked for verification.", 9, False, GREY, True)
st.merge_cells("B15:K15"); st.row_dimensions[15].height = 28
req_rows = [[r["project"], r["variant"], r["top"], r["fp"], r["stage"], r["code"], r["desc"], r["qty"], r.get("source", ""), r.get("verified", "")] for r in REQ]
make_table(st, "tblReq", 17, 2, REQ_COLS, req_rows, widths=[13, 17, 16, 9, 9, 16, 40, 9, 46, 18])
section(st, "N4", "Kitting stages (report column order)")
make_table(st, "tblStages", 5, 14, STAGE_COLS, [[s["stage"], s["order"], s["desc"]] for s in STAGES], widths=[11, 8, 30])
label(st, "N14", "To add a stage later (e.g. Internal Unit, Housing): add a row here with its order number, then add its assembly lines to the requirements table.", 9, False, GREY, True)
st.merge_cells("N14:P16")
st["N14"].alignment = WRAP

# ---------------------------------------------------------------- paste sheets
def write_paste(ws, rows, fmts, note):
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            if v is None: continue
            c = ws.cell(i + 1, j + 1, v)
            if fmts and i >= 1 and j < len(fmts[min(len(fmts) - 1, 1)]):
                f = fmts[min(len(fmts) - 1, 1)][j]
                if f and f != "General": c.number_format = f
    for j in range(len(rows[0]) if rows else 0):
        ws.cell(1, j + 1).font = font(10, True)
    ws.freeze_panes = "A2"
write_paste(asp, ASP, ASP_FMT, "")
write_paste(co, COOIS, COOIS_FMT, "")
for c in "ABCDEFGHIJKLMNO": co.column_dimensions[c].width = 15
co.column_dimensions["D"].width = 38; co.column_dimensions["O"].width = 34
asp.column_dimensions["D"].width = 18; asp.column_dimensions["H"].width = 18

# ---------------------------------------------------------------- Kitting Log
title(lg, "Kitting Log", "Every confirmed Kitting Progress entry and the COOIS orders allocated to it. Paste new rows from Record Kitting at the bottom. "
      "Never delete a row: to cancel an entry type Y in Void on any of its rows and give a reason.")
calc_log = {
    "Order key": 'TRIM([@Order]&"")',
    "Active": 'IF([@[Entry ID]]="",IF(UPPER(TRIM([@Void]&""))="Y",0,1),IF(COUNTIFS([Entry ID],[@[Entry ID]],[Void],"Y")>0,0,1))',
}
chk = let(("cid", '[@[Entry ID]]&""'), ("cty", 'UPPER(TRIM([@[Record type]]&""))'), ("cv", 'TRIM([@[Variant ID]]&"")'), ("cs", 'TRIM([@Stage]&"")'),
          ("ca", 'TRIM([@[Assembly code]]&"")'), ("ck", '[@[Order key]]'),
          ("cnv", 'COUNTIF(tblReq[Variant ID],cv)'), ("cns", 'COUNTIFS(tblReq[Variant ID],cv,tblReq[Stage],cs)'),
          ("cna", 'COUNTIFS(tblReq[Variant ID],cv,tblReq[Stage],cs,tblReq[Assembly code],ca)'),
          ("cnva", 'COUNTIFS(tblReq[Variant ID],cv,tblReq[Assembly code],ca)'),
          ("cdup", 'COUNTIFS([Order key],ck,[Record type],"COOIS order",[Active],1)'),
          ("cri", 'XMATCH(ck,Rv_Key)'),
          ("crm", 'IF(ISNUMBER(cri),INDEX(Rv_Mat,cri)&"","")'), ("crq", 'IF(ISNUMBER(cri),N(INDEX(Rv_Qty,cri)),0)'),
          ("crc", 'IF(ISNUMBER(cri),INDEX(Rv_CLSD,cri)&"","")'),
          ("creps", 'N([@[REPs kitted]])'), ("cf", '[@[REP from]]'), ("ct", 'IF(ISNUMBER([@[REP to]]),[@[REP to]],[@[REP from]])'),
          ("casp", 'COUNTIFS(P_Var,cv,P_Valid,1)>0'),
          'IF(cid="","",IF([@Active]=0,"Void",'
          'IF(cty="COOIS ORDER",IF(ck="","Order number missing",IF(cdup>1,"Duplicate allocation - fix",IF(ISNA(cri),"Not in current COOIS - review",'
          'IF(UPPER(crm)<>UPPER(ca),"Material changed - review",IF(crc="Yes","CLSD - review",IF(ABS(crq-N([@[Order qty]]))>0.000000001,"Qty changed - review",'
          'IF(cnva=0,"Assembly not set up for variant","OK"))))))),'
          'IF(cty="KITTING PROGRESS",IF(cnv=0,"Assembly setup required",IF(cns=0,"Stage not set up for variant",IF(AND(ca<>"",cna=0),"Assembly not in this stage",'
          'IF(NOT(creps>0),"REPs kitted missing",IF(ISNUMBER(cf),IF(ct<cf,"REP range invalid",IF(AND(ca="",ABS(ct-cf+1-creps)>0.000000001),"REP range differs from REPs kitted",'
          'IF(casp,"OK","Variant not in current ASP - review"))),IF(casp,"OK","Variant not in current ASP - review")))))),'
          '"Record type not recognised"))))')
calc_log["Check"] = chk
log_rows = [[r.get(k, None) if r.get(k, "") != "" else None for k in LOG_COLS] for r in LOGROWS]
make_table(lg, "tblLog", 5, 1, LOG_COLS, log_rows, calc_log, style="TableStyleLight9",
           widths=[10, 11, 15, 16, 9, 16, 9, 8, 8, 12, 8, 6, 22, 26, 12, 7, 30])
for r in range(6, 6 + max(1, len(log_rows)) + 200):
    lg.cell(r, 2).number_format = "dd-mmm-yy"
lg.freeze_panes = "A6"
label(lg, "A4", "Columns A-N are typed or pasted. Grey columns (Order key, Active, Check) calculate automatically. Record type is \"Kitting Progress\" or \"COOIS order\".", 9, False, GREY, True)

# ---------------------------------------------------------------- COOIS Review
title(cr, "COOIS Review", "Read-only view of the pasted COOIS export: SAP status tokens beside the tracker's allocation. CNF alone is never treated as kitted.")
cr_cols = ["Order", "Material Number", "Material description", "Order qty", "Release date (actual)", "Delivered qty", "System Status",
           "PCNF", "CNF", "DLV", "CLSD", "TECO", "Settings rows using code", "Allocated entry", "Allocated variant", "Review"]
hdr(cr, 5, 1, cr_cols, INFO)
def tok(t): return f'IF(ISNUMBER(SEARCH(" {t} ",cst)),"Yes","")'
coois_rv = let(("craw", f"COOIS!$A$2:$O${COOIS_ROWS}"), ("ckz", 'TRIM(INDEX(craw,,1)&"")'), ("ckeep", 'ckz<>""'),
               ("cr", "FILTER(craw,ckeep)"), ("ckk", "FILTER(ckz,ckeep)"), ("cmat", 'TRIM(INDEX(cr,,3)&"")'),
               ("cst", '" "&UPPER(TRIM(INDEX(cr,,15)&""))&" "'),
               ("cnv", "COUNTIF(tblReq[Assembly code],cmat)"),
               ("clk", 'TRIM(tblLog[Order]&"")'), ("clm", '(tblLog[Record type]="COOIS order")*(tblLog[Active]=1)*(clk<>"")'),
               ("cae", 'IFERROR(XLOOKUP(ckk,FILTER(clk,clm),FILTER(tblLog[Entry ID]&"",clm),""),"")'),
               ("cav", 'IFERROR(XLOOKUP(ckk,FILTER(clk,clm),FILTER(tblLog[Variant ID]&"",clm),""),"")'),
               ("caa", 'IFERROR(XLOOKUP(ckk,FILTER(clk,clm),FILTER(TRIM(tblLog[Assembly code]&""),clm),""),"")'),
               ("ccl", tok("CLSD")),
               ("crev", 'IF(ccl="Yes","CLSD - review",IF(cnv=0,"Code not in Settings",IF(cae="",IF(cnv>1,"Shared code - no contract allocation","Not allocated"),'
                        'IF(UPPER(caa)<>UPPER(cmat),"Material changed - review","Allocated"))))'),
               'HSTACK(ckk,cmat,INDEX(cr,,4)&"",IFERROR(--INDEX(cr,,7),0),IF(INDEX(cr,,9)="","",INDEX(cr,,9)),IFERROR(--INDEX(cr,,10),0),TRIM(INDEX(cr,,15)&""),'
               + ",".join(tok(t) for t in ["PCNF", "CNF", "DLV", "CLSD", "TECO"]) + ',cnv,cae,cav,crev)')
n_coois = max(1, sum(1 for r in COOIS[1:] if r and r[0] not in (None, "")))
da(cr, "A6", coois_rv, f"A6:P{5 + n_coois}")
for j, w in enumerate([12, 16, 36, 8, 12, 9, 30, 6, 6, 6, 6, 6, 10, 10, 16, 32]): cr.column_dimensions[CL(1 + j)].width = w
for c in "E": pass
for r in range(6, 6 + n_coois + 50): cr.cell(r, 5).number_format = "dd-mmm-yy"
cr.freeze_panes = "B6"
for nm, col in [("Rv_Key", "A"), ("Rv_Mat", "B"), ("Rv_Qty", "D"), ("Rv_Rel", "E"), ("Rv_CLSD", "K"), ("Rv_Entry", "N"), ("Rv_Review", "P")]:
    name(nm, f"'COOIS Review'!${col}$6:${col}${RV_LAST}")
label(cr, "B4", "CNF can be recorded before physical picking, so the tracker never reads CNF as kitted. CLSD is shown for review; its cause is not confirmed by the status alone.", 9, False, GREY, True)
fx(cr, "H2", 'COUNTIFS(Rv_Review,"Not allocated")+COUNTIFS(Rv_Review,"Shared code - no contract allocation")')
fx(cr, "H3", 'COUNTIFS(Rv_Review,"CLSD - review")+COUNTIFS(Rv_Review,"Material changed - review")')
label(cr, "I2", "orders for configured codes not yet allocated"); label(cr, "I3", "orders needing review (CLSD / material changed)")
cr["H2"].font = font(12, True, NAVY); cr["H3"].font = font(12, True, CRIT)

# ---------------------------------------------------------------- Calc
calc["A1"] = "Calculations - nothing to type on this sheet. Every result here can be traced to the ASP, COOIS, Settings or Kitting Log."
calc["A1"].font = font(11, True, GREY)
label(calc, "P1", "ASP header"); label(calc, "P2", "ASP week Monday"); label(calc, "P3", "Kit week Monday"); label(calc, "P4", "Kit week")
LC = ASP_LASTCOL
da(calc, "Q1", f'IF(ASP!$Q$1:${LC}$1="","",ASP!$Q$1:${LC}$1&"")', f"Q1:{LC}1")
da(calc, "Q2", let(("wh", f"TRIM(ASP!$Q$1:${LC}$1&\"\")"), ("wy", '2000+VALUE(LEFT(wh,FIND(" ",wh)-1))'),
                   ("ww", 'VALUE(MID(wh,FIND("W",UPPER(wh))+1,3))'), ("wj", "DATE(wy,1,4)"),
                   'IFERROR(wj-WEEKDAY(wj,2)+1+7*(ww-1),"")'), f"Q2:{LC}2")
da(calc, "Q3", let(("wm", f"IFERROR(Q2:{LC}2+0,-1)"), 'IF(wm<0,"",wm-7*KitLead)'), f"Q3:{LC}3")
da(calc, "Q4", let(("wd", f"IFERROR(Q3:{LC}3+0,-1)"), 'IF(wd<0,"",RIGHT(YEAR(wd+3),2)&" W"&ISOWEEKNUM(wd))'), f"Q4:{LC}4")
name("WkHdr", f"Calc!$Q$1:${LC}$1"); name("WkASPMon", f"Calc!$Q$2:${LC}$2"); name("WkKitMon", f"Calc!$Q$3:${LC}$3"); name("WkKitLbl", f"Calc!$Q$4:${LC}$4")
name("AspWeeks", f"ASP!$Q$2:${LC}${ASP_ROWS}")

plan_cols = ["ASP row", "Project Id", "Project Name", "Product", "Planning code", "Prod Id", "Variant ID", "Seabed Temp", "Body Type", "Nber FP",
             "System Qty", "Planned Qty", "Prod. progress", "DR2M calc", "DR2M demand", "Weekly plan total", "Due to kit by end of week",
             "Due to kit before this week", "Due to kit next 4 weeks",
             "Row in use", "Valid variant", "Requirement lines", "Built before plan", "Complete kits", "Partly kitted", "Remaining to kit",
             "Progress", "Due this week", "Behind", "Units without COOIS order", "Top code differs", "Status"]
hdr(calc, 7, 1, plan_cols, "8A94A6")
PE = 7 + NPLAN   # last plan row (8..1006)
da(calc, "A8", let(("pinfo", f"ASP!$C$2:$P${ASP_ROWS}"), ("pblk", "IFERROR(--AspWeeks,0)"), ("pkm", "WkKitMon"),
                   ("pone", "SEQUENCE(COLUMNS(pblk),1,1,0)"),
                   'HSTACK(SEQUENCE(ROWS(pinfo),1,2),IF(pinfo="","",pinfo),MMULT(pblk,pone),MMULT(pblk,TRANSPOSE(--(pkm<=ThisMon))),'
                   'MMULT(pblk,TRANSPOSE(--(pkm<ThisMon))),MMULT(pblk,TRANSPOSE((pkm>ThisMon)*(pkm<=ThisMon+28))))'), f"A8:S{PE}")
def pc(col): return f"Calc!${col}$8:${col}${PE}"
for nm, col in [("P_ProjId", "B"), ("P_ProjName", "C"), ("P_Product", "D"), ("P_ProdId", "F"), ("P_Var", "G"), ("P_System", "K"),
                ("P_Planned", "L"), ("P_WSum", "P"), ("P_DueEnd", "Q"), ("P_DueBefore", "R"), ("P_Next4", "S"), ("P_RowOK", "T"),
                ("P_Valid", "U"), ("P_Lines", "V"), ("P_Built", "W"), ("P_Complete", "X"), ("P_Partial", "Y"), ("P_Remaining", "Z"),
                ("P_Pct", "AA"), ("P_DueThis", "AB"), ("P_Behind", "AC"), ("P_Short", "AD"), ("P_TopDiff", "AE"), ("P_Status", "AF")]:
    name(nm, pc(col))
da(calc, "T8", let(("qfr", f'TRIM($G$8:$G${PE}&"")'), ("qpr", f'$D$8:$D${PE}&""'),
                   ("qok", f'($B$8:$B${PE}&$C$8:$C${PE}&qpr&qfr&$F$8:$F${PE})<>""'),
                   ("qva", '(qfr<>"")*(UPPER(qfr)<>"NA")*(UPPER(qfr)<>"N/A")'),
                   ("qnl", "IF(qva,COUNTIF(tblReq[Variant ID],qfr),0)"),
                   ("qsy", f"IFERROR(--$K$8:$K${PE},0)"), ("qpl", f"IFERROR(--$L$8:$L${PE},0)"),
                   ("qcm", "MINIFS(L_CovReps,L_Var,qfr)"), ("qcz", "IF(qnl>0,qcm,0)"),
                   ("qpa", "MAXIFS(L_StartReps,L_Var,qfr)-qcz"),
                   ("qde", f"$Q$8:$Q${PE}"), ("qdb", f"$R$8:$R${PE}"), ("qmx", "IF(qcz>qdb,qcz,qdb)"),
                   ("qdu", "IF(qde-qmx>0,qde-qmx,0)"), ("qbe", "IF(qdb-qcz>0,qdb-qcz,0)"),
                   ("qtp", 'IFERROR(XLOOKUP(qfr,tblReq[Variant ID],TRIM(tblReq[Top code]&"")),"")'),
                   ("qtd", f'(qnl>0)*(UPPER(qtp)<>UPPER(TRIM($F$8:$F${PE}&"")))'),
                   ("qst", 'IF(qok=0,"",IF(qva=0,"No variant ID in ASP",IF(qnl=0,"Assembly setup required",IF(qtd,"Top code differs - review",'
                           'IF(qpl<=0,"No plan quantity",IF(qcz>=qpl,"All kitted",IF(qbe>0,"Kit behind",IF(qdu>0,"Kit due this week","On track"))))))))'),
                   'HSTACK(--qok,qva,qnl,qsy-qpl,IF(qnl>0,qcm,""),IF(qnl>0,qpa,""),IF(qnl>0,IF(qpl-qcz>0,qpl-qcz,0),""),'
                   'IF((qnl>0)*(qpl>0),qcz/IF(qpl=0,1,qpl),"not available"),IF(qnl>0,qdu,""),IF(qnl>0,qbe,""),'
                   'IF(qnl>0,SUMIFS(L_Short,L_Var,qfr),""),--qtd,qst)'), f"T8:AF{PE}")

line_cols = ["Variant ID", "Stage", "Stage order", "Assembly code", "Assembly description", "Qty per REP", "Recorded units",
             "Allocated units", "Covered units", "Units without COOIS order", "Allocated beyond recorded", "Complete REPs (line)",
             "Started REPs (line)", "Planned REPs", "Units still to kit", "Unallocated COOIS qty", "Settings rows using code"]
hdr(calc, 7, 34, line_cols, "8A94A6")
n_req = max(1, len(REQ))
LE = 7 + LINE_CAP
da(calc, "AH8", let(("lv", 'TRIM(tblReq[Variant ID]&"")'), ("ls", 'TRIM(tblReq[Stage]&"")'), ("lc", 'TRIM(tblReq[Assembly code]&"")'),
                    ("lq", "IFERROR(--tblReq[Qty per REP],0)"), ("lok", 'lv<>""'),
                    ("lso", 'IFERROR(XLOOKUP(ls,tblStages[Stage],tblStages[Order]),"")'),
                    ("lrw", 'SUMIFS(tblLog[REPs kitted],tblLog[Record type],"Kitting Progress",tblLog[Variant ID],lv,tblLog[Stage],ls,tblLog[Assembly code],"",tblLog[Active],1)'),
                    ("lrl", 'SUMIFS(tblLog[REPs kitted],tblLog[Record type],"Kitting Progress",tblLog[Variant ID],lv,tblLog[Stage],ls,tblLog[Assembly code],lc,tblLog[Active],1)'),
                    ("lrec", "(lrw+lrl)*lq"),
                    ("lal", 'SUMIFS(tblLog[Order qty],tblLog[Record type],"COOIS order",tblLog[Variant ID],lv,tblLog[Assembly code],lc,tblLog[Active],1)'),
                    ("lcov", "IF(lrec<lal,lrec,lal)"),
                    ("lcr", "IF(lq>0,INT(lcov/IF(lq=0,1,lq)+0.000000001),0)"), ("lsr", "IF(lq>0,-INT(-lcov/IF(lq=0,1,lq)+0.000000001),0)"),
                    ("lpl", "SUMIFS(P_Planned,P_Var,lv)"),
                    ("lav", 'SUMIFS(Rv_Qty,Rv_Mat,lc,Rv_Entry,"",Rv_CLSD,"")'),
                    'FILTER(EXPAND(HSTACK(lv,ls,lso,lc,tblReq[Assembly description]&"",lq,lrec,lal,lcov,lrec-lcov,lal-lcov,lcr,lsr,lpl,'
                    'IF(lpl*lq-lcov>0,lpl*lq-lcov,0),lav,COUNTIF(tblReq[Assembly code],lc)),ROWS(lv)+1,,""),VSTACK(lok,FALSE),"")'), f"AH8:AX{7 + n_req}")
def lc_(col): return f"Calc!${col}$8:${col}${LE}"
for nm, col in [("L_Var", "AH"), ("L_Stage", "AI"), ("L_Code", "AK"), ("L_Rec", "AN"), ("L_Alloc", "AO"), ("L_Cov", "AP"), ("L_Short", "AQ"),
                ("L_CovReps", "AS"), ("L_StartReps", "AT"), ("L_Avail", "AW")]:
    name(nm, lc_(col))
# lists for dropdowns
label(calc, "BA6", "Lists for dropdowns", 10, True)
hdr(calc, 7, 53, ["Products", "Projects", "Variants (all)", "Configured variants", "Stages for entry", "Codes for entry"], "8A94A6")
da(calc, "BA8", 'IFERROR(SORT(UNIQUE(FILTER(P_Product,(P_Product<>"")*(P_RowOK=1)))),"REPEATERS")', "BA8:BA40")
da(calc, "BB8", 'IFERROR(VSTACK("All",SORT(UNIQUE(FILTER(P_ProjName,(UPPER(P_Product)=UPPER(Dashboard!$C$5))*(P_ProjName<>""))))),"All")', "BB8:BB60")
da(calc, "BC8", 'IFERROR(VSTACK("All",SORT(UNIQUE(FILTER(P_Var,(UPPER(P_Product)=UPPER(Dashboard!$C$5))*(P_Valid=1)*'
                 '((Dashboard!$F$5="All")+(Dashboard!$F$5="")+(P_ProjName=Dashboard!$F$5)>0))))),"All")', "BC8:BC120")
da(calc, "BD8", 'IFERROR(SORT(UNIQUE(FILTER(TRIM(tblReq[Variant ID]&""),TRIM(tblReq[Variant ID]&"")<>""))),"(none)")', "BD8:BD60")
da(calc, "BE8", let(("es", "SORTBY(tblStages[Stage],tblStages[Order])"),
                    'FILTER(es,COUNTIFS(tblReq[Variant ID],\'Record Kitting\'!$C$7,tblReq[Stage],es)>0,es)'), "BE8:BE30")
da(calc, "BF8", 'VSTACK("",FILTER(tblReq[Assembly code],(tblReq[Variant ID]=\'Record Kitting\'!$C$7)*(tblReq[Stage]=\'Record Kitting\'!$C$8),""))', "BF8:BF40")
name("ProductList", "Calc!$BA$8:$BA$200"); name("ProjectList", "Calc!$BB$8:$BB$400"); name("VariantAllList", "Calc!$BC$8:$BC$600")
name("VariantList", "Calc!$BD$8:$BD$400"); name("StageList", "Calc!$BE$8:$BE$100"); name("CodeList", "Calc!$BF$8:$BF$200")
for col in range(1, 60): calc.column_dimensions[CL(col)].width = 12

# ---------------------------------------------------------------- Record Kitting
title(rec, "Record Kitting", "Step 1 describe what was physically kitted.  Step 2 review the quantities.  Step 3 copy the green block into the Kitting Log.")
section(rec, "B5", "Step 1 - Kitting Progress entry")
inputs = [("Entry date (blank = report date)", "dd-mmm-yy"), ("Variant ID (contract variant)", "@"), ("Stage", "@"),
          ("Assembly code (optional - blank = whole stage)", "@"), ("REPs kitted (repeaters' worth)", "0.##"),
          ("REP from (optional)", "0"), ("REP to (optional)", "0"), ("Note (e.g. management instruction)", "@")]
for i, (lab, fmt) in enumerate(inputs):
    r = 6 + i; label(rec, f"B{r}", lab); c = rec[f"C{r}"]; c.fill = fill(INPUT); c.border = box; c.number_format = fmt; c.protection = UNLOCK
sc_in = sc.get("entry", {})
for i, k in enumerate(["date", "variant", "stage", "assembly", "reps", "from", "to", "note"]):
    if sc_in.get(k) not in (None, ""): rec[f"C{6 + i}"] = sc_in[k]
rec.column_dimensions["B"].width = 40; rec.column_dimensions["C"].width = 22
for c, w in zip("DEFGHIJKL", [40, 8, 9, 10, 10, 11, 10, 10, 11]): rec.column_dimensions[c].width = w
section(rec, "E5", "Checks")
RK_V, RK_S, RK_A, RK_R, RK_F, RK_T = "$C$7", "$C$8", "$C$9", "$C$10", "$C$11", "$C$12"
common = [("kv", f'TRIM({RK_V}&"")'), ("ks", f'TRIM({RK_S}&"")'), ("ka", f'TRIM({RK_A}&"")'), ("kr", f"N({RK_R})"),
          ("kf", RK_F), ("kt", f"IF(ISNUMBER({RK_T}),{RK_T},{RK_F})"),
          ("km", '(TRIM(tblReq[Variant ID]&"")=kv)*(TRIM(tblReq[Stage]&"")=ks)*(((ka="")+(TRIM(tblReq[Assembly code]&"")=ka))>0)'),
          ("knv", "COUNTIF(tblReq[Variant ID],kv)"), ("kn", "SUM(km)"),
          ("kcodes", 'FILTER(TRIM(tblReq[Assembly code]&""),km,"")'), ("kq", "FILTER(IFERROR(--tblReq[Qty per REP],0),km,0)"),
          ("kst", 'FILTER(TRIM(tblReq[Stage]&""),km,"")'),
          ("krec", "SUMIFS(L_Rec,L_Var,kv,L_Stage,kst,L_Code,kcodes)"), ("kal", "SUMIFS(L_Alloc,L_Var,kv,L_Stage,kst,L_Code,kcodes)"),
          ("kneed", "IF(krec+kr*kq-kal>0,krec+kr*kq-kal,0)")]
da(rec, "E6", let(*common,
                  ("kmsg", 'VSTACK(IF(kv="","Choose a variant.",""),IF(AND(kv<>"",knv=0),"Assembly setup required for this variant (Settings).",""),'
                           'IF(AND(knv>0,kn=0),"This stage / assembly code is not set up for this variant.",""),IF(kr<0,"REPs kitted cannot be negative.",""),'
                           'IF(AND(ISNUMBER(kf),kt<kf),"REP to is before REP from.",""),'
                           'IF(AND(ISNUMBER(kf),kt>=kf,ka="",kr>0,ABS(kt-kf+1-kr)>0.000000001),"Warning: REP range covers "&(kt-kf+1)&" repeaters but REPs kitted is "&kr&".",""),'
                           'IF(AND(kv<>"",COUNTIFS(P_Var,kv,P_Valid,1)=0),"Warning: variant not in the current ASP - entry allowed, flagged for review.",""),'
                           'IF(AND(kn>0,kr=0),"No REPs kitted entered - only orders for earlier shortages will be proposed.",""),'
                           'IF(COUNTIF($K$19:$K$200,">0")>0,"Warning: not enough unallocated COOIS orders - see Shortage after. The shortfall stays open.",""))'),
                  'FILTER(kmsg,kmsg<>"","No problems found")'), "E6:E14")
rec["E6"].alignment = WRAP
section(rec, "B16", "Step 2 - review quantities for this entry")
rv_cols = ["Assembly code", "Description", "Qty per REP", "Units this entry", "Already recorded", "Already allocated",
           "Needs COOIS orders", "Orders available", "Orders proposed", "Shortage after", "Settings rows using code"]
hdr(rec, 18, 2, rv_cols, TEAL)
da(rec, "B19", let(*common, ("kprop", 'SUMIFS($X$19:$X$5000,$P$19:$P$5000,"COOIS order",$S$19:$S$5000,kcodes)'),
                   ("kav", "SUMIFS(L_Avail,L_Var,kv,L_Stage,kst,L_Code,kcodes)"),
                   'IF(kn=0,"",HSTACK(kcodes,FILTER(tblReq[Assembly description]&"",km,""),kq,kr*kq,krec,kal,kneed,kav,kprop,'
                   'IF(kneed-kprop>0,kneed-kprop,0),COUNTIF(tblReq[Assembly code],kcodes)))'), "B19:L30")
section(rec, "N16", "Step 3 - copy the green block (from N19, all rows) and Paste Values at the bottom of the Kitting Log")
hdr(rec, 18, 14, LOG_COLS, OK)
da(rec, "N19", let(*common, ("kdt", "IF(ISNUMBER($C$6),$C$6,RptDate)"), ("knote", '$C$13&""'),
                   ("kblock", "OR(kv=\"\",knv=0,kn=0,kr<0,AND(ISNUMBER(kf),kt<kf))"),
                   ("rkey", 'IFERROR(Rv_Key&"","")'), ("rmat", 'IFERROR(Rv_Mat&"","")'), ("rqty", "IFERROR(--Rv_Qty,0)"),
                   ("rrel", "IFERROR(--Rv_Rel,0)"), ("rent", 'IFERROR(Rv_Entry&"","x")'), ("rcl", 'IFERROR(Rv_CLSD&"","Yes")'),
                   ("cmk", 'ISNUMBER(XMATCH(rmat,kcodes))*(rent="")*(rcl<>"Yes")*(rkey<>"")'),
                   ("hasc", "SUM(cmk)>0"),
                   ("cand", "SORTBY(FILTER(HSTACK(rmat,rrel,rkey,rqty),cmk),FILTER(UPPER(rmat),cmk),1,FILTER(rrel,cmk),1,FILTER(rkey,cmk),1)"),
                   ("cmt", "INDEX(cand,,1)"), ("cqq", "INDEX(cand,,4)"), ("cn", "ROWS(cand)"),
                   ("ccum", "MMULT((SEQUENCE(cn)>=SEQUENCE(1,cn))*(cmt=TRANSPOSE(cmt)),cqq)"),
                   ("cnd", "XLOOKUP(cmt,kcodes,kneed,0)"),
                   ("pick", 'IF(hasc,FILTER(EXPAND(cand,cn+1,,""),VSTACK((ccum-cqq)<cnd,FALSE),""),"")'),
                   ("np", 'IF(ISERROR(INDEX(pick,1,1)),0,IF(INDEX(pick,1,1)="",0,ROWS(pick)))'),
                   ("kid", '"K"&TEXT(1+MAX(0,IFERROR(--MID(FILTER(tblLog[Entry ID]&"",LEFT(tblLog[Entry ID]&"",1)="K"),2,9),0)),"00000")'),
                   ("prow", 'HSTACK(kid,kdt,"Kitting Progress",kv,ks,ka,kr,IF(ISNUMBER(kf),kf,""),IF(ISNUMBER(' + RK_T + '),' + RK_T + ',""),"","","","",knote)'),
                   ("kone", "SEQUENCE(MAX(np,1))^0"),
                   ("orow", 'HSTACK(IF(kone,kid),IF(kone,kdt),IF(kone,"COOIS order"),IF(kone,kv),IF(kone,ks),INDEX(pick,,1),IF(kone,""),IF(kone,""),IF(kone,""),'
                            'INDEX(pick,,3),INDEX(pick,,4),IF(kone,""),IF(kone,""),IF(kone,""))'),
                   'IF(kblock,"Nothing to paste - see Checks",IF(kr>0,IF(np>0,VSTACK(prow,orow),prow),IF(np>0,orow,"Nothing to paste - no REPs kitted and no orders needed")))'),
   "N19:AA80")
for c, w in zip(range(14, 28), [9, 10, 15, 16, 8, 16, 7, 7, 7, 12, 7, 5, 8, 20]): rec.column_dimensions[CL(c)].width = w
for r in range(19, 400): rec.cell(r, 15).number_format = "dd-mmm-yy"
rec.freeze_panes = "A6"
label(rec, "B17", "Orders are proposed oldest release date first, then lowest order number, skipping allocated and CLSD orders. A shortfall stays open: "
      "when a later COOIS has more orders, enter the same variant and stage with REPs kitted = 0.", 8, False, GREY, True)
rec.merge_cells("B17:L17"); rec.row_dimensions[17].height = 24; rec["B17"].alignment = WRAP
label(rec, "B4", "No macros: confirming = copy the green block (Step 3) and Paste Special > Values on the first empty row of the Kitting Log table. "
      "Excel cannot append the rows for you or stop a second paste; the Kitting Log Check column flags duplicates.", 9, True, CRIT, False)
rec.merge_cells("B4:L4"); rec.row_dimensions[4].height = 26; rec["B4"].alignment = WRAP

# ---------------------------------------------------------------- Dashboard
title(dash, "Kitting Progress Dashboard")
fx(dash, "B3", '"Week "&RIGHT(YEAR(ThisMon+3),2)&" W"&ISOWEEKNUM(ThisMon)&"  |  report date "&TEXT(RptDate,"dd mmm yyyy")&"  |  kit target = ASP week - "&KitLead&" weeks  |  ASP: "&ASP!$B$2&" (export "&IFERROR(TEXT(ASP!$A$2,"dd mmm yyyy"),"?")&")  |  COOIS: "&Settings!$C$10&IF(Settings!$C$11="",""," loaded "&TEXT(Settings!$C$11,"dd mmm yyyy"))')
dash["B3"].font = font(10, False, GREY, True)
for cell, lab in [("B5", "Product"), ("E5", "Project"), ("H5", "Variant")]:
    label(dash, cell, lab, 10, True); dash[cell].alignment = Alignment(horizontal="right", vertical="center")
for cell, val in [("C5", SETTINGS.get("defaultProduct", "REPEATERS")), ("F5", "All"), ("I5", "All")]:
    dash[cell] = sc.get("dash", {}).get(cell, val); dash[cell].fill = fill(INPUT); dash[cell].border = box; dash[cell].protection = UNLOCK; dash[cell].font = font(11, True)
dash.merge_cells("C5:D5"); dash.merge_cells("F5:G5"); dash.merge_cells("I5:K5")
# scope helper (hidden column AD)
SM = ('(UPPER(P_Product)=UPPER($C$5))*(P_RowOK=1)*((($F$5="All")+($F$5=""))+(UPPER(P_ProjName)=UPPER($F$5))>0)'
      '*((($I$5="All")+($I$5=""))+(UPPER(P_Var)=UPPER($I$5))>0)')
SCF = SM + "*(P_Lines>0)"
kpis = {"AD2": f"SUMPRODUCT({SM},IFERROR(--P_Planned,0))", "AD3": f"SUMPRODUCT({SCF},IFERROR(--P_Planned,0))",
        "AD4": f"SUMPRODUCT({SCF},IFERROR(--P_Complete,0))", "AD5": f"SUMPRODUCT({SCF},IFERROR(--P_Partial,0))",
        "AD6": f"SUMPRODUCT({SCF},IFERROR(--P_DueThis,0))", "AD7": f"SUMPRODUCT({SCF},IFERROR(--P_Behind,0))",
        "AD8": f"SUMPRODUCT({SCF})", "AD9": f"SUMPRODUCT({SM}*(P_Lines=0)*(IFERROR(--P_Planned,0)>0))",
        "AD10": f"SUMPRODUCT({SCF},IFERROR(--P_Short,0))",
        "AD11": "'COOIS Review'!$H$2", "AD12": "'COOIS Review'!$H$3+COUNTIFS(tblLog[Check],\"*review*\")+COUNTIFS(tblLog[Check],\"*fix*\")+COUNTIFS(tblLog[Check],\"Assembly*\")+COUNTIFS(tblLog[Check],\"Stage*\")+COUNTIFS(tblLog[Check],\"REP*\")+COUNTIFS(tblLog[Check],\"Order number missing\")+COUNTIFS(tblLog[Check],\"Record type*\")"}
for c, f in kpis.items(): da(dash, c, f); dash[c].font = font(8, False, "9AA1AE")
dash.column_dimensions["AD"].hidden = True
tiles = [("B", "C", "PLANNED", '=IF(AD8=0,AD2,AD3)', '=IF(AD8=0,"ASP planned qty in scope (no variant set up)","ASP planned qty, set-up variants"&IF(AD2>AD3," ("&AD2&" incl. not set up)",""))'),
         ("D", "E", "COMPLETE ASSEMBLY KITS", "=AD4", '="REPs with every set-up line covered"'),
         ("F", "G", "PARTLY KITTED", "=AD5", '="REPs started, not complete"'),
         ("H", "I", "REMAINING TO KIT", "=IF(AD8=0,\"n/a\",MAX(0,AD3-AD4))", '="planned minus complete kits"'),
         ("J", "K", "KITTING PROGRESS", '=IF(OR(AD8=0,AD3=0),"not available",AD4/AD3)', '="complete kits / planned (set-up variants)"'),
         ("L", "M", "KITS DUE THIS WEEK", "=AD6", '="kit target this week, not yet complete"'),
         ("N", "O", "KITS BEHIND", "=AD7", '="kit target before this week, not complete"'),
         ("P", "Q", "COOIS: UNALLOCATED / REVIEW", '=AD11&" / "&AD12', '="orders, all contracts - see COOIS Review"')]
for a, b, lab, val, sub in tiles:
    dash.merge_cells(f"{a}7:{b}7"); dash.merge_cells(f"{a}8:{b}8"); dash.merge_cells(f"{a}9:{b}9")
    dash[f"{a}7"] = lab; dash[f"{a}7"].font = font(8, True, GREY); dash[f"{a}7"].alignment = CEN
    dash[f"{a}8"] = val; dash[f"{a}8"].font = font(20, True, NAVY); dash[f"{a}8"].alignment = CEN
    dash[f"{a}9"] = sub; dash[f"{a}9"].font = font(8, False, GREY, True); dash[f"{a}9"].alignment = CEN
    for r in (7, 8, 9):
        for col in (a, b): dash[f"{col}{r}"].fill = fill(PANEL)
dash["J8"].number_format = "0%"; dash["N8"].font = font(20, True, CRIT); dash["L8"].font = font(20, True, WARN)
dash.row_dimensions[8].height = 34; dash.row_dimensions[9].height = 26
section(dash, "B11", "Kits due by kit week (scope above, set-up variants)  -  this week highlighted")
for i, t in enumerate(["Kit week", "REPs due to kit that week", "Cumulative due", "Complete kits to date", "Short of cumulative due"]):
    label(dash, f"B{12 + i}", t, 9, True)
da(dash, "C12", let(("smk", SCF), ("sw", "MMULT(TRANSPOSE(smk),IFERROR(--AspWeeks,0))"), ("swk", "TRANSPOSE(WkKitMon)"),
                    ("sM", "ThisMon+7*SEQUENCE(1,13,-4)"), ("sdue", "MMULT(sw,--(swk=sM))"), ("scum", "MMULT(sw,--(swk<=sM))"),
                    'VSTACK(RIGHT(YEAR(sM+3),2)&" W"&ISOWEEKNUM(sM),sdue,scum,sM*0+$AD$4,IF(scum-$AD$4>0,scum-$AD$4,0))'), "C12:O16")
section(dash, "B18", "Data checks")
checks = [
    ('IF(AND(ASP!$C$1="Project Id",ASP!$D$1="Project Name",ASP!$E$1="Product",ASP!$G$1="Prod Id",ASP!$H$1="Friendly_name",ASP!$K$1="Nber FP",ASP!$L$1="System Qty",ASP!$M$1="Planned Qty"),"OK  ASP columns as expected ("&COUNTIFS(P_RowOK,1)&" ASP lines, "&COUNTIFS(WkHdr,"?? W*")&" week columns)","CHECK  ASP headers are not in the expected columns - paste the export at ASP!A1")'),
    ('IF(SUMPRODUCT(--(COOIS!$A$1:$O$1=({' + ",".join(f'"{h}"' for h in ["Order", "Order Type", "Material Number", "Material description", "Plant", "Batch", "Order quantity (GMEIN)", "Unit of measure (=GMEIN)", "Release date (actual)", "Delivered quantity (GMEIN)", "Created on", "Entered by", "Change date", "Last changed by", "System Status"]) + '})))=15,"OK  COOIS columns Order..System Status in the expected order ("&COUNTA(Rv_Key)&" orders)","CHECK  COOIS headings differ from the expected layout - paste the export at COOIS!A1")'),
    ('LET(nl,IFERROR(Calc!$BD$8:$BD$400&"",""),nm,(nl<>"")*(nl<>"(none)")*(COUNTIFS(P_Var,nl,P_Valid,1)=0),IF(SUM(nm)=0,"OK  every set-up variant is in the current ASP","REVIEW  set up but not in current ASP (history kept): "&TEXTJOIN(", ",TRUE,IF(nm=1,nl,""))))'),
    ('IF(SUMPRODUCT((P_Lines>0)*(COUNTIFS(P_Var,P_Var,P_Valid,1)>1))=0,"OK  no set-up variant appears on more than one ASP line","REVIEW  a set-up variant appears on more than one ASP line - quantities are added together")'),
    ('IF(AD12=0,"OK  no Kitting Log or COOIS rows need review",AD12&" rows need review - see Check column in Kitting Log and Review column in COOIS Review")'),
    ('IF(AD10=0,"OK  every recorded unit in scope is covered by a COOIS order",AD10&" recorded units in scope have no allocated COOIS order yet (not counted as complete)")'),
    ('IF(AD9=0,"OK  every planned line in scope has assembly requirements",AD9&" planned ASP lines in scope have no requirements - Assembly setup required")')]
for i, f in enumerate(checks):
    da(dash, f"B{19 + i}", f); dash[f"B{19 + i}"].font = font(9); dash[f"B{19 + i}"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=False)
section(dash, "B27", "Contract overview")
ov_cols = ["Project", "Variant ID", "Top code", "FP", "Body", "Temp", "System qty", "Built before plan", "Planned", "Kits due by this week",
           "Complete kits", "Partly kitted", "Remaining", "Progress", "Due this week", "Behind", "Units without COOIS order", "Status"]
hdr(dash, 28, 2, ov_cols)
da(dash, "B29", let(("om", SM), ("orows", f"FILTER(Calc!$A$8:$AF${PE},om,\"\")"),
                    'IF(SUM(om)=0,"No ASP lines for this selection",CHOOSECOLS(orows,3,7,6,10,9,8,11,23,12,17,24,25,26,27,28,29,30,32))'), "B29:S90")
for c, w in zip("BCDEFGHIJKLMNOPQRS", [16, 17, 15, 6, 7, 7, 8, 9, 8, 9, 9, 8, 9, 9, 9, 8, 10, 22]): dash.column_dimensions[c].width = w
for r in range(29, 400): dash[f"O{r}"].number_format = "0%"
dash.freeze_panes = "A7"
defs = ["Definitions",
        "Planned: ASP Planned Qty for the selected product/project/variant. Repeaters and amplifiers are separate products in the ASP and are never added together.",
        "Complete assembly kits: repeaters' worth for which every set-up assembly line is recorded as kitted AND covered by allocated COOIS orders (the lowest line decides).",
        "Partly kitted: repeaters with at least one line started but not all lines complete.  Remaining: planned minus complete kits.",
        "Kit target week = ASP planned week minus the kit lead (Settings, 3 weeks). Due this week / behind compare complete kits with the cumulative plan, contract by contract.",
        "A complete assembly kit is not a finished repeater. COOIS CNF is not treated as kitted. Contract allocation is the controller's entry, not a link proven by SAP."]
# put definitions to the right of KPI tiles instead (stable position)
for i, t in enumerate(defs):
    c = dash.cell(5 + i, 20, t); c.font = font(9, i == 0, GREY if i else NAVY); c.alignment = WRAP
dash.column_dimensions["T"].width = 70

# ---------------------------------------------------------------- Repeater View
title(rv, "Repeater View", "Repeater-by-repeater stage progress for one variant. Kitted = recorded against that REP number. Kitted* = covered by a contract-level quantity, "
      "shown against the earliest open REP positions; it is NOT a physical traceability link.")
label(rv, "B5", "Variant", 10, True); rv["C5"] = sc.get("rvVariant", "E2A-R12P-T5-C"); rv["C5"].fill = fill(INPUT); rv["C5"].border = box; rv["C5"].protection = UNLOCK
rv["C5"].font = font(11, True); rv.merge_cells("C5:E5")
V5 = "$C$5"
facts = [("Project", f'IFERROR(INDEX(P_ProjName,XMATCH(1,(UPPER(P_Var)=UPPER({V5}))*(P_Valid=1))),"not in ASP")'),
         ("Top code (ASP)", f'IFERROR(INDEX(P_ProdId,XMATCH(1,(UPPER(P_Var)=UPPER({V5}))*(P_Valid=1))),"")'),
         ("System qty", f"SUMIFS(P_System,P_Var,{V5})"), ("Built before plan", f"SUMIFS(P_System,P_Var,{V5})-SUMIFS(P_Planned,P_Var,{V5})"),
         ("Planned", f"SUMIFS(P_Planned,P_Var,{V5})"), ("Weekly plan total", f"SUMIFS(P_WSum,P_Var,{V5})"),
         ("Complete kits", f'IF(COUNTIF(tblReq[Variant ID],{V5})=0,"Assembly setup required",MINIFS(L_CovReps,L_Var,{V5}))'),
         ("This week", 'RIGHT(YEAR(ThisMon+3),2)&" W"&ISOWEEKNUM(ThisMon)')]
for i, (lab, f) in enumerate(facts):
    col = CL(7 + i)
    rv[f"{col}4"] = lab; rv[f"{col}4"].font = font(8, True, GREY); rv[f"{col}4"].alignment = CEN; rv[f"{col}4"].fill = fill(PANEL)
    da(rv, f"{col}5", f); rv[f"{col}5"].font = font(12, True, NAVY); rv[f"{col}5"].alignment = CEN; rv[f"{col}5"].fill = fill(PANEL)
fx(rv, "G6", f'IF(SUMIFS(P_WSum,P_Var,{V5})<>SUMIFS(P_Planned,P_Var,{V5}),"Note: the weekly plan total differs from Planned Qty in the ASP - REPs beyond the weekly plan show \'Not in weekly plan\'.","")')
rv["G6"].font = font(9, False, WARN, True)
rvc = [("rv", f'TRIM({V5}&"")'), ("rmk", '(UPPER(TRIM(P_Var&""))=UPPER(rv))*(rv<>"")'),
       ("rsy", "SUMPRODUCT(rmk,IFERROR(--P_System,0))"), ("rpl", "SUMPRODUCT(rmk,IFERROR(--P_Planned,0))"), ("rb", "rsy-rpl"),
       ("res", "SORTBY(tblStages[Stage],tblStages[Order])"),
       ("rnst", "COUNTIFS(tblReq[Variant ID],rv,tblReq[Stage],res)"),
       ("rstg", "TRANSPOSE(FILTER(res,rnst>0,\"\"))"), ("rS", "SUM(--(rnst>0))"),
       ("rcm", "MINIFS(L_CovReps,L_Var,rv,L_Stage,rstg)"), ("rsa", "MAXIFS(L_StartReps,L_Var,rv,L_Stage,rstg)"),
       ("rsh", "SUMIFS(L_Short,L_Var,rv,L_Stage,rstg)"),
       ("rn", "SEQUENCE(MAX(rsy,1))"),
       ("rem", '(tblLog[Active]=1)*(UPPER(TRIM(tblLog[Record type]&""))="KITTING PROGRESS")*(UPPER(TRIM(tblLog[Variant ID]&""))=UPPER(rv))'
               '*(TRIM(tblLog[Assembly code]&"")="")*ISNUMBER(tblLog[REP from])*(IFERROR(--tblLog[REP from],0)>0)'),
       ("rhe", "SUM(rem)>0"),
       ("ref", "FILTER(tblLog[REP from],rem,0)"), ("ret", "FILTER(IF(ISNUMBER(tblLog[REP to]),tblLog[REP to],tblLog[REP from]),rem,0)"),
       ("rstx", 'FILTER(UPPER(TRIM(tblLog[Stage]&"")),rem,"")'),
       ("rA", "IF(rhe,MMULT((rn>=TRANSPOSE(ref))*(rn<=TRANSPOSE(ret)),--(rstx=UPPER(rstg)))>0,(rn+SEQUENCE(1,rS))<0)"),
       ("rasg", "MMULT(TRANSPOSE(rn^0),--rA)"), ("rpool", "IF(rcm-rasg>0,rcm-rasg,0)")]
section(rv, "B8", "Stage summary (contract level)")
for i, t in enumerate(["Stage", "Complete kits (covered)", "Recorded against REP numbers", "Contract level, not assigned to a REP",
                       "Partly kitted REPs", "Units without COOIS order"]):
    label(rv, f"B{9 + i}", t, 9, i == 0); rv.merge_cells(f"B{9 + i}:D{9 + i}")
da(rv, "E9", let(*rvc, 'IF(rS=0,"Assembly setup required",VSTACK(rstg,rcm,rasg,rpool,rsa-rcm,rsh))'), "E9:L14")
section(rv, "B16", "Repeater by repeater")
da(rv, "B17", let(("rv", f'TRIM({V5}&"")'), ("res", "SORTBY(tblStages[Stage],tblStages[Order])"),
                  ("rnst", "COUNTIFS(tblReq[Variant ID],rv,tblReq[Stage],res)"),
                  'IF(SUM(--(rnst>0))=0,"",HSTACK("REP","ASP week",TRANSPOSE(FILTER(res,rnst>0,"")),"Next kit to issue","Kit due week","Weeks left","Status"))'), "B17:M17")
da(rv, "B18", let(*rvc,
                  ("relig", "(rA=FALSE)*(rn>rb)"), ("rrk", "MMULT(--(rn>=TRANSPOSE(rn)),relig)"),
                  ("rcell", 'IF(rn<=rb,"Built",IF(rA,"Kitted",IF(rrk<=rpool,"Kitted*",IF(rrk<=rpool+rsa-rcm,"Part",""))))'),
                  ("rnk", '(rcell<>"Kitted")*(rcell<>"Kitted*")'),
                  ("rsc", "MMULT(rnk*2^(rS-SEQUENCE(1,rS)),SEQUENCE(rS)^0)"),
                  ("rfi", "IF(rsc=0,0,rS-INT(LOG(IF(rsc=0,1,rsc),2)+0.000000001))"),
                  ("rnx", 'IF(rn<=rb,"",IF(rfi=0,"All kitted",INDEX(rstg,1,IF(rfi=0,1,rfi))))'),
                  ("rwk", "MMULT(TRANSPOSE(rmk),IFERROR(--AspWeeks,0))"), ("rnw", "COLUMNS(rwk)"),
                  ("rcum", "MMULT(rwk,--(SEQUENCE(rnw)<=SEQUENCE(1,rnw)))"),
                  ("ridx", "IFERROR(XMATCH(rn-rb,rcum,1),0)"),
                  ("rkm", 'IF(ridx=0,"",INDEX(WkKitMon,1,IF(ridx=0,1,ridx)))'),
                  ("rwl", 'IF(rkm="","",ROUND((rkm-ThisMon)/7,0))'),
                  ("rstt", 'IF(rn<=rb,"Built before plan",IF(rfi=0,"All kitted",IF(rkm="","Not in weekly plan",IF(rwl<0,"Kit behind",IF(rwl=0,"Kit due this week","On track")))))'),
                  'IF(rS=0,"Assembly setup required for this variant",IF(rsy<=0,"No repeaters for this variant in the current ASP",'
                  'HSTACK(rn,IF((rn<=rb)+(ridx=0),"",INDEX(WkHdr,1,IF(ridx=0,1,ridx))),rcell,rnx,IF((rn<=rb)+(ridx=0),"",INDEX(WkKitLbl,1,IF(ridx=0,1,ridx))),'
                  'IF(rn<=rb,"",rwl),rstt)))'), "B18:M400")
for c, w in zip("BCDEFGHIJKLM", [8, 9, 9, 9, 9, 9, 9, 16, 11, 9, 16, 12]): rv.column_dimensions[c].width = w
for c in "GHIJKLMN": rv.column_dimensions[c].width = max(rv.column_dimensions[c].width or 0, 13)
rv.freeze_panes = "C18"

# ---------------------------------------------------------------- conditional formats
def cf_text(ws, rng, text, fg, bg, exact=True):
    tl = rng.split(":")[0]
    f = f'{tl}="{text}"' if exact else f'ISNUMBER(SEARCH("{text}",{tl}))'
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f], font=Font(color=fg, bold=True), fill=fill(bg)))
for rng in ["B18:M400"]:
    cf_text(rv, rng, "Kitted", OK, OK_S); cf_text(rv, rng, "Kitted*", OK, OK_L); cf_text(rv, rng, "Part", WARN, WARN_S)
    cf_text(rv, rng, "Built", GREY, CALC); cf_text(rv, rng, "Kit behind", CRIT, CRIT_S); cf_text(rv, rng, "Kit due this week", WARN, WARN_S)
    cf_text(rv, rng, "On track", INFO, INFO_S); cf_text(rv, rng, "All kitted", OK, OK_S); cf_text(rv, rng, "Not in weekly plan", WARN, WARN_S)
for t, fg, bg in [("Kit behind", CRIT, CRIT_S), ("Kit due this week", WARN, WARN_S), ("On track", INFO, INFO_S), ("All kitted", OK, OK_S),
                  ("Assembly setup required", GREY, CALC), ("Top code differs - review", CRIT, CRIT_S), ("No plan quantity", GREY, CALC)]:
    cf_text(dash, "S29:S400", t, fg, bg)
dash.conditional_formatting.add("C12:O16", FormulaRule(formula=[X("C$12=RIGHT(YEAR(ThisMon+3),2)&\" W\"&ISOWEEKNUM(ThisMon)")], fill=fill(WARN_S), font=Font(bold=True)))
cf_text(dash, "B19:B25", "OK", OK, WHITE, exact=False)
for t in ["CHECK", "REVIEW", "rows need", "recorded units", "planned ASP"]:
    dash.conditional_formatting.add("B19:B25", FormulaRule(formula=[f'LEFT(B19,{len(t)})="{t}"'], font=Font(color=CRIT, bold=True)))
lg.conditional_formatting.add("Q6:Q20000", FormulaRule(formula=['AND(Q6<>"",Q6<>"OK",Q6<>"Void")'], font=Font(color=CRIT, bold=True), fill=fill(CRIT_S)))
lg.conditional_formatting.add("A6:N20000", FormulaRule(formula=['$P6=0'], font=Font(color="9AA1AE", strike=True)))
for t, fg, bg in [("CLSD - review", CRIT, CRIT_S), ("Material changed - review", CRIT, CRIT_S), ("Shared code - no contract allocation", WARN, WARN_S),
                  ("Allocated", OK, OK_L)]:
    cf_text(cr, "P6:P20005", t, fg, bg)
rec.conditional_formatting.add("N19:AA2000", FormulaRule(formula=['N19<>""'], fill=fill(PASTE)))
rec.conditional_formatting.add("E6:E14", FormulaRule(formula=['OR(LEFT(E6,7)="Warning",LEFT(E6,2)="No")'], font=Font(color=WARN, bold=True)))
rec.conditional_formatting.add("E6:E14", FormulaRule(formula=['AND(E6<>"",E6<>"No problems found",LEFT(E6,7)<>"Warning",LEFT(E6,2)<>"No")'], font=Font(color=CRIT, bold=True)))
rec.conditional_formatting.add("K19:K200", FormulaRule(formula=['N(K19)>0'], font=Font(color=CRIT, bold=True), fill=fill(CRIT_S)))

# ---------------------------------------------------------------- data validation
def dv_list(ws, cell, src, prompt=None):
    d = DataValidation(type="list", formula1=src, allow_blank=True, showErrorMessage=False)
    if prompt: d.promptTitle = "Choose"; d.prompt = prompt; d.showInputMessage = True
    ws.add_data_validation(d); d.add(cell)
dv_list(dash, "C5", "ProductList"); dv_list(dash, "F5", "ProjectList"); dv_list(dash, "I5", "VariantAllList")
dv_list(rv, "C5", "VariantList"); dv_list(rec, "C7", "VariantList"); dv_list(rec, "C8", "StageList"); dv_list(rec, "C9", "CodeList")
d = DataValidation(type="list", formula1='"Kitting Progress,COOIS order"', allow_blank=True); lg.add_data_validation(d); d.add("C6:C20000")
d = DataValidation(type="list", formula1='"Y"', allow_blank=True); lg.add_data_validation(d); d.add("L6:L20000")

# ---------------------------------------------------------------- Read Me
readme.sheet_view.showGridLines = False; readme.column_dimensions["A"].width = 2; readme.column_dimensions["B"].width = 4; readme.column_dimensions["C"].width = 120
readme["B2"] = "ASN Kitting Progress Tracker - Read Me"; readme["B2"].font = font(18, True, NAVY)
readme["B3"] = "For Production Controllers. No macros. Yellow cells and the two Settings tables are inputs; ASP and COOIS are paste areas; everything else calculates."
readme["B3"].font = font(10, False, GREY, True)
RM = [
 ("1. Load a new ASP", ["Open the ASP export, select all (Ctrl+A), copy, then click cell A1 on the ASP sheet and paste. Do not insert, delete or move columns.",
   "Type the file name and today's date in Settings (ASP file loaded / ASP loaded on). The export's own version and export time show at the top of the Dashboard.",
   "New contracts and variants appear in the Dashboard filters automatically. Recorded progress is kept by Variant ID (ASP Friendly_name), so it is never erased by a new ASP.",
   "If a variant you have set up or recorded is no longer in the ASP, the Dashboard Data checks list it for review. Its history stays in the Kitting Log."]),
 ("2. Load COOIS (keep the column order)", ["Run COOIS for the assembly codes you need, export to Excel, copy everything and paste at cell A1 of the COOIS sheet. Columns Order to System Status must stay in SAP's order; extra columns to the right are fine.",
   "Order numbers and material codes are treated as text identifiers. An order already in an earlier export keeps its allocation; nothing is counted twice.",
   "An allocated order that is missing from a new export is flagged 'Not in current COOIS - review'. That does not prove it was cancelled - the new export may cover a different selection.",
   "The COOIS Review sheet shows each order's SAP status tokens (PCNF, CNF, DLV, CLSD, TECO are matched as whole words, so PCNF is never read as CNF) beside the tracker's allocation."]),
 ("3. Record what was kitted (daily routine)", ["Record Kitting sheet. Step 1: choose the Variant ID you actually worked on (even if management changed the order), the Stage, and how many repeaters' worth you kitted. Leave Assembly code blank for a whole stage, or choose one code for a partial (line-only) entry.",
   "REP from / REP to are optional. Fill them only when you know which repeater numbers the kit is for. Without them the quantity is held at contract level.",
   "Step 2 shows, for each assembly line: units for this entry, what is already recorded and allocated, units still needing COOIS orders, orders available and any shortage.",
   "Step 3 is the green block. Select it from N19 to the last green row, copy, go to the Kitting Log, click column A on the first empty row of the table and use Paste Special > Values. That is the confirmation. Then clear the Step 1 inputs.",
   "Excel without macros cannot press a Confirm button for you or stop you pasting twice. The Check column catches a second paste (Duplicate allocation - fix): void the second entry."]),
 ("4. Shortages and shared assemblies", ["If fewer suitable COOIS orders exist than the quantity needed, the shortfall shows in Step 2 and in 'Units without COOIS order'. The stage is not counted as complete until it is covered.",
   "When a later COOIS export contains the missing orders, enter the same variant and stage with REPs kitted = 0. Only the outstanding orders are proposed.",
   "Codes used by more than one variant (for example the PSB 91YAF03404AAA) are never assigned to a contract automatically. Unallocated orders show 'Shared code - no contract allocation'."]),
 ("5. Add or change a contract's assembly requirements", ["Settings > requirements table: one row per assembly line with Project, Variant ID (exactly as the ASP Friendly_name), Top code, Fibre pairs, Stage, Assembly code, description and Qty per REP (for ONE repeater).",
   "Type the new rows in the first empty row below the table; it grows by itself. The same code can appear in several variants with different quantities.",
   "Fill 'Verified by / date' once a line is checked against the official kitting sheet. A variant without rows shows 'Assembly setup required'.",
   "New stages (Internal Unit, Housing...) are added in the Kitting stages table with an order number, then their lines in the requirements table."]),
 ("6. Correct an entry", ["Never delete or overwrite log rows. In the Kitting Log type Y in Void on any row of the wrong entry and write a Void reason. The whole entry (progress and its orders) stops counting, stays visible and is struck through.",
   "Then record the correct entry through Record Kitting as normal. The released orders become available again.",
   "To assign contract-level kitting to known repeater numbers later: void the original entry and re-enter it with REP from / REP to."]),
 ("7. Save, back up and hand over", ["Save as .xlsx (no macros). Keep a dated copy each week, e.g. 'ASN Kitting Tracker 2026-W44.xlsx', in the shared drive.",
   "The Kitting Log and Settings are the only records that cannot be re-created from SAP. Everything else is recalculated from the pasted ASP and COOIS.",
   "Handover: give the new controller this file and this Read Me. Nothing is hidden: the Calc sheet shows every intermediate number.",
   "If the file opens in Protected View, click Enable Editing - formulas do not calculate until then."]),
 ("8. What the terms mean", ["Kitting Progress: what the controller records as physically kitted (repeaters' worth of a stage or line).",
   "Status: the result shown in the report - All kitted, On track, Kit due this week, Kit behind, Assembly setup required, review flags.",
   "Complete assembly kits: repeaters' worth where every set-up assembly line is recorded AND covered by allocated COOIS orders. It is not a finished repeater.",
   "Kitted (Repeater View): recorded against that REP number. Kitted*: covered by contract-level quantity and shown against the earliest open REP positions - not a physical traceability link. Part: some lines started.",
   "Kit target week: ASP planned week minus the kit lead (3 weeks). Weeks crossing a year end (including 53-week years) are handled with ISO dates.",
   "Kits behind: repeaters whose kit target week is before this week, minus complete kits (contract by contract). Due this week: target this week, not yet complete.",
   "CNF: can be recorded before physical picking, so it never counts as kitted. CLSD: shown for review; the cause (e.g. failed build or test) is not confirmed by the status alone.",
   "Contract allocation is the controller's entry. The COOIS export does not prove which contract an order belongs to."])]
r = 5
for head, items in RM:
    readme.cell(r, 2, head).font = font(12, True, TEAL); r += 1
    for t in items:
        readme.cell(r, 2, "-").font = font(10, False, GREY); c = readme.cell(r, 3, t); c.font = font(10); c.alignment = WRAP
        readme.row_dimensions[r].height = 15 * (1 + len(t) // 140); r += 1
    r += 1
readme.cell(r, 2, "Sheets").font = font(12, True, TEAL); r += 1
for t in ["Dashboard - KPIs, kits due by week, data checks and contract overview for the chosen product / project / variant.",
          "Repeater View - stage-by-stage progress for every REP of one variant, next kit to issue and kit due week.",
          "Record Kitting - the entry form (Steps 1-3).  Kitting Log - the permanent record.  Settings - requirements, stages, kit lead, report date.",
          "COOIS Review - the pasted orders with SAP status tokens and allocation.  ASP / COOIS - paste areas.  Calc - workings, read only."]:
    readme.cell(r, 2, "-"); c = readme.cell(r, 3, t); c.font = font(10); c.alignment = WRAP; r += 1

# ---------------------------------------------------------------- protection (no password) for report sheets
for ws in (dash, rv, cr, calc, readme, rec):
    ws.protection.sheet = True; ws.protection.formatCells = False; ws.protection.formatColumns = False
    ws.protection.autoFilter = False; ws.protection.sort = False
for ws in wb.worksheets:
    for sh, cell, _ in DA:
        if sh == ws.title:
            for mr in ws.merged_cells.ranges:
                assert cell not in mr, f"dynamic array {sh}!{cell} inside merged range {mr}"
wb.active = 1
wb.save(OUT)

# ---------------------------------------------------------------- dynamic array metadata (Excel spill) / test widening
def finish(path):
    z = zipfile.ZipFile(path); files = {n: z.read(n) for n in z.namelist()}; order = z.namelist(); z.close()
    meta = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<metadata xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:xda="http://schemas.microsoft.com/office/spreadsheetml/2017/dynamicarray">'
            '<metadataTypes count="1"><metadataType name="XLDAPR" minSupportedVersion="120000" copy="1" pasteAll="1" pasteValues="1" merge="1" splitFirst="1" rowColShift="1" clearFormats="1" clearComments="1" assign="1" coerce="1" cellMeta="1"/></metadataTypes>'
            '<futureMetadata name="XLDAPR" count="1"><bk><extLst><ext uri="{bdbb8cdc-fa1e-496e-a857-3c3f30c029c3}"><xda:dynamicArrayProperties fDynamic="1" fCollapsed="0"/></ext></extLst></bk></futureMetadata>'
            '<cellMetadata count="1"><bk><rc t="1" v="0"/></bk></cellMetadata></metadata>').encode()
    files["xl/metadata.xml"] = meta
    if "xl/metadata.xml" not in order: order.append("xl/metadata.xml")
    ct = files["[Content_Types].xml"].decode()
    if "metadata.xml" not in ct:
        ct = ct.replace("</Types>", '<Override PartName="/xl/metadata.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheetMetadata+xml"/></Types>')
    files["[Content_Types].xml"] = ct.encode()
    rel = files["xl/_rels/workbook.xml.rels"].decode()
    if "sheetMetadata" not in rel:
        rel = rel.replace("</Relationships>", '<Relationship Id="rIdDAmeta" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sheetMetadata" Target="metadata.xml"/></Relationships>')
    files["xl/_rels/workbook.xml.rels"] = rel.encode()
    # sheet xml file per title
    wbx = files["xl/workbook.xml"].decode(); rels = files["xl/_rels/workbook.xml.rels"].decode()
    sheetfile = {}
    for m in re.finditer(r'<sheet [^>]*name="([^"]+)"[^>]*r:id="([^"]+)"', wbx):
        t = re.search(r'Id="' + m.group(2) + r'"[^>]*Target="/?(?:xl/)?([^"]+)"', rels) or re.search(r'Target="/?(?:xl/)?([^"]+)"[^>]*Id="' + m.group(2) + '"', rels)
        sheetfile[m.group(1).replace("&amp;", "&")] = "xl/" + t.group(1)
    widen = {}
    over = sc.get("_widen", {})
    for sh, cell, tref in DA:
        tref = over.get(f"{sh}!{cell}", tref)
        if TEST and tref: widen.setdefault(sheetfile[sh], {})[cell] = tref
    for n in list(files):
        if n.startswith("xl/worksheets/sheet") and n.endswith(".xml"):
            x = files[n].decode()
            x = re.sub(r'<c r="([A-Z]+\d+)"((?:(?! cm=)[^>])*)><f t="array"', lambda m: f'<c r="{m.group(1)}"{m.group(2)} cm="1"><f t="array"', x)
            for cell, tref in widen.get(n, {}).items():
                x = re.sub(r'(<c r="' + cell + r'"[^>]*><f t="array" ref=")' + cell + '"', lambda m: m.group(1) + tref + '"', x)
            files[n] = x.encode()
    z = zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED)
    for n in order: z.writestr(n, files[n])
    z.close()
finish(OUT)
print("built", OUT, "test" if TEST else "", "DA cells:", len(DA))
